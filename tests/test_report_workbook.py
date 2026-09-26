"""Save report's tables in one Excel workbook: a sheet per table as its own
export writes it, the statistics under it, names Excel accepts, a Contents
sheet, and charts left out."""

from __future__ import annotations

from pathlib import Path

import matplotlib
import numpy as np
import openpyxl
import pandas as pd
import pytest

from siamang.data import SurveyData
from siamang.reporting import Report
from siamang.reporting.workbook import SheetNames, save_tables

matplotlib.use("Agg")

DOCUMENTS = Path(__file__).resolve().parent / "documents"


@pytest.fixture(scope="module")
def questionnaire_doc():
    from siamang.model import loads

    return loads((DOCUMENTS / "brand_awareness.questionnaire.json").read_text("utf-8"))


@pytest.fixture(scope="module")
def survey(questionnaire_doc):
    from siamang.model import from_document

    return from_document(questionnaire_doc).survey


@pytest.fixture(scope="module")
def data(survey) -> SurveyData:
    return survey.simulate(n=200, seed=3)


def _rows(sheet) -> list[tuple]:
    return [row for row in sheet.iter_rows(values_only=True)]


def _own_export(table, tmp_path) -> dict[str, list[tuple]]:
    file = table.export_xlsx(tmp_path / f"own_{id(table)}.xlsx")
    book = openpyxl.load_workbook(file)
    return {sheet.title: _rows(sheet) for sheet in book.worksheets}


def test_sheet_names_are_ones_excel_accepts():
    names = SheetNames(reserved=["Contents"])
    long = "Satisfaction with the service in the shop, by region [weighted]"
    first = names.take(long)
    assert first == "Satisfaction with the service i" and len(first) == 31
    assert names.take(long) == "Satisfaction with the servi (2)"
    assert names.take("Q1: which/brand?*") == "Q1 which brand"
    assert names.take("contents") == "contents (2)"  # the Contents sheet's name is taken
    assert names.take("History") == "History (2)"  # Excel keeps History for itself
    assert names.take("'quoted'") == "quoted"
    with_suffix = names.take(long, suffix="Post-hoc")
    assert with_suffix == "Satisfaction with th – Post-hoc" and len(with_suffix) == 31
    assert names.take("AGE") == "AGE" and names.take("age") == "age (2)"


def test_every_table_gets_its_sheet_as_its_own_export_writes_it(data, tmp_path):
    freq = data.report.freq("region")
    means = data.report.means("age", by="region", method="anova", posthoc="tukey")
    banner = data.report.banner(["satisfaction"], ["region", "gender"])
    coefficients = pd.DataFrame({"term": ["const", "age"], "coef": [2.5, 0.0125]}, index=["a", "b"])
    report = (
        Report(title="Brand study")
        .heading("Who answered")
        .add(freq, caption="Region of residence")
        .add(data.plot.bar("region"), caption="A chart")  # not a table
        .add({"N": 200}, caption="Statistics")  # not a table
        .heading("Age")
        .add(means)
        .add(banner, caption="Satisfaction by region and gender")
        .add(coefficients)
    )
    path = report.save_tables(tmp_path / "out" / "report.xlsx")
    assert path.exists()
    book = openpyxl.load_workbook(path)
    assert book.sheetnames == [
        "Contents",
        "Region of residence",
        "Age",
        "Age – Post-hoc",
        "Satisfaction by region and gend",
        "Age (2)",
    ]

    # Each table as its own export_xlsx writes it, its statistics under it.
    own = _own_export(freq, tmp_path)["Table"]
    rows = _rows(book["Region of residence"])
    assert rows[: len(own)] == own
    assert rows[len(own)] == (None,) * len(own[0])
    assert rows[len(own) + 1][:2] == ("Variable", "Region")
    assert rows[len(own) + 2][:2] == ("N valid", 200)

    exported = _own_export(means, tmp_path)
    assert set(exported) == {"Table", "Post-hoc"}
    assert _rows(book["Age"])[: len(exported["Table"])] == exported["Table"]
    pairs = _rows(book["Age – Post-hoc"])
    assert pairs[: len(exported["Post-hoc"])] == exported["Post-hoc"]
    assert pairs[1][0] == "Capital vs North"
    # The pairs' own statistics under them, as the report prints them: which
    # way a difference runs, and that p is adjusted already.
    footer = {row[0]: row[1] for row in pairs[len(exported["Post-hoc"]) + 1 :]}
    assert footer == {
        key: value if isinstance(value, int | float | str) else str(value)
        for key, value in means.posthoc_table.stats.items()
    }
    assert footer["Method"] == "Tukey HSD" and footer["Difference"].startswith("mean of the first")
    assert footer["p"].startswith("adjusted for the number of pairs")
    stats = {row[0]: row[1] for row in _rows(book["Age"])[len(exported["Table"]) + 1 :]}
    assert stats["Test"] == "One-way ANOVA" and stats["Post-hoc"].startswith("Tukey HSD")
    assert isinstance(stats["p"], float) and isinstance(stats["N"], int)

    # The banner keeps its significance letters.
    letters = _rows(book["Satisfaction by region and gend"])
    assert letters[: len(banner.to_frame()) + 1] == _own_export(banner, tmp_path)["Table"]
    assert any(isinstance(cell, str) and cell.endswith(" C") for row in letters for cell in row)

    # A bare DataFrame is written as the report prints it: no index.
    assert _rows(book["Age (2)"]) == [("term", "coef"), ("const", 2.5), ("age", 0.0125)]

    contents = _rows(book["Contents"])
    assert contents[0][0] == "Brand study"
    assert contents[3] == ("Sheet", "Section", "Table")
    assert contents[4:] == [
        ("Region of residence", "Who answered", "Region of residence"),
        ("Age", "Age", "Group means: Age"),
        ("Age – Post-hoc", "Age", "Group means: Age — Post-hoc"),
        ("Satisfaction by region and gend", "Age", "Satisfaction by region and gender"),
        ("Age (2)", "Age", "Table"),
    ]
    link = book["Contents"]["A6"]
    assert link.hyperlink.location == "'Age'!A1" or link.hyperlink.target == "#'Age'!A1"


def test_a_report_without_tables_says_so(tmp_path):
    path = save_tables(Report(title="Only words").text("Nothing to count."), tmp_path / "r.xlsx")
    book = openpyxl.load_workbook(path)
    assert book.sheetnames == ["Contents"]
    assert book["Contents"]["A4"].value == "This report has no tables."


def test_the_tables_go_to_an_xlsx_workbook_only(tmp_path):
    with pytest.raises(ValueError, match="written to an .xlsx workbook; got 'r.xls'"):
        Report().save_tables(tmp_path / "r.xls")


def test_a_frame_with_two_header_rows_gets_one(tmp_path):
    frame = pd.DataFrame(
        [[1, 2]], columns=pd.MultiIndex.from_tuples([("Region", "North"), ("Region", "South")])
    )
    book = openpyxl.load_workbook(
        Report().add(frame, caption="Wide").save_tables(tmp_path / "w.xlsx")
    )
    assert _rows(book["Wide"]) == [("Region / North", "Region / South"), (1, 2)]


# ─── the node ────────────────────────────────────────────────────────────────


def _flow(save_params):
    nodes = [
        ("src", "source.responses", {}),
        ("xtab", "analyze.crosstab", {"row": "gender", "col": "region"}),
        (
            "means",
            "analyze.means",
            {"y": "age", "by": "region", "method": "anova", "posthoc": "tukey"},
        ),
        ("bar", "visualize.bar", {"variable": "region", "show": "percent"}),
        (
            "section",
            "output.report_section",
            {"heading": "Sample", "captions": {"xtab.table": "Gender by region"}},
        ),
        ("save", "output.save_report", {"title": "Brand study", **save_params}),
    ]
    edges = [("src", "data", n, "data") for n in ("xtab", "means", "bar")]
    edges += [
        ("xtab", "table", "section", "items"),
        ("means", "table", "section", "items"),
        ("bar", "chart", "section", "items"),
        ("section", "report", "save", "sections"),
    ]
    return {
        "schema_version": "1.0",
        "name": "t",
        "nodes": [{"id": i, "type": t, "params": p} for i, t, p in nodes],
        "edges": [
            {"from": {"node": a, "port": ap}, "to": {"node": b, "port": bp}}
            for a, ap, b, bp in edges
        ],
    }


def test_a_stored_save_report_renders_the_code_it_always_did(questionnaire_doc):
    from siamang.flow.document import resolve_flow
    from siamang.flow.template import render_node

    graph = resolve_flow(_flow({}), questionnaire=questionnaire_doc)
    assert render_node(graph, "save") == (
        "n_save = Report.combine([n_section], title='Brand study', toc=False, theme=None)\n"
        'n_save.provenance(os.environ.get("SIAMANG_PROVENANCE"))\n'
        "n_save.save('outputs/report.md')\n"
        "n_save.save(Path('outputs/report.md').with_suffix(\".html\"))\n"
    )


def test_save_report_writes_the_workbook_beside_the_report(questionnaire_doc, survey, tmp_path):
    from siamang.flow import FlowRunner, check_flow, generate_flow

    flow = _flow({"xlsx": True, "path": "outputs/brand.md"})
    assert check_flow(flow, questionnaire=questionnaire_doc) == []
    code = generate_flow(flow, questionnaire_doc)
    assert 'n_save.save_tables(Path("outputs/brand.md").with_suffix(".xlsx"))' in code
    simulated = survey.simulate(n=150, seed=2)
    frame = simulated.frame.assign(w=np.random.default_rng(1).uniform(0.5, 2, 150))
    data = SurveyData(frame=frame, variables=simulated.variables, questionnaire=survey)
    FlowRunner(flow, questionnaire=survey).run(sources={"src": data}, cwd=tmp_path)
    book = openpyxl.load_workbook(tmp_path / "outputs" / "brand.xlsx")
    assert book.sheetnames == ["Contents", "Gender by region", "Sample", "Sample – Post-hoc"]
    assert (tmp_path / "outputs" / "brand.md").exists()

    FlowRunner(_flow({}), questionnaire=survey).run(sources={"src": data}, cwd=tmp_path / "plain")
    assert not list((tmp_path / "plain").rglob("*.xlsx"))


def test_text_that_begins_with_an_equals_sign_is_written_as_text(data, tmp_path):
    """openpyxl stores a string beginning with "=" as a formula: an open answer
    =HYPERLINK(…) became a live link in the workbook, and read back as nothing."""
    from siamang.io import read_snapshot, write_snapshot

    frame = data.frame.copy()
    frame["comment"] = frame["comment"].astype(object)
    frame.loc[frame.index[:5], "comment"] = '=HYPERLINK("http://evil.example","Click me")'
    frame.loc[frame.index[5:10], "comment"] = "=1+1"
    answers = data.with_frame(frame)
    report = (
        Report(title="=cmd|' /C calc'!A0")
        .heading("=SUM(A1:A9)")
        .add(answers.report.freq("comment"), caption="=1+2")
        .add(pd.DataFrame({"said": ["=2*3", "fine"]}))
    )
    book = openpyxl.load_workbook(report.save_tables(tmp_path / "r.xlsx"))
    cells = [cell for sheet in book.worksheets for row in sheet.iter_rows() for cell in row]
    assert not [cell.coordinate for cell in cells if cell.data_type == "f"]
    texts = {cell.value for cell in cells if isinstance(cell.value, str)}
    assert {'=HYPERLINK("http://evil.example","Click me")', "=1+1", "=2*3", "=1+2"} <= texts
    assert book["Contents"]["A1"].value == "=cmd|' /C calc'!A0"
    assert book["Contents"]["B5"].value == "=SUM(A1:A9)"
    # A table's own export, and the data written to Excel, keep the text too.
    own = openpyxl.load_workbook(answers.report.freq("comment").export_xlsx(tmp_path / "t.xlsx"))
    assert not [c for row in own.active.iter_rows() for c in row if c.data_type == "f"]
    back = read_snapshot(write_snapshot(answers, tmp_path / "data.xlsx"))
    assert list(back.frame["comment"][:6]) == list(frame["comment"][:6])
    assert (back.frame["comment"] == "=1+1").sum() == 5


def test_a_link_to_a_sheet_whose_name_has_an_apostrophe_leads_to_it(data, tmp_path):
    """Excel names a sheet in a reference quoted, an apostrophe inside doubled:
    'Brand's image'!A1 leads nowhere, 'Brand''s image'!A1 to the sheet."""
    from siamang.io.excel_text import sheet_link

    assert sheet_link("Brand's image") == "'Brand''s image'!A1"
    report = Report(title="Brands").add(data.report.freq("region"), caption="Brand's image")
    book = openpyxl.load_workbook(report.save_tables(tmp_path / "r.xlsx"))
    assert book.sheetnames == ["Contents", "Brand's image"]
    link = book["Contents"]["A5"]
    assert link.value == "Brand's image"
    assert link.hyperlink.target == "#'Brand''s image'!A1"


def test_the_contents_name_the_tables_of_the_later_analyses(tmp_path):
    """Without a caption, a Perceptual map's three tables in one section were
    'Table', 'Table' and 'Table'; each now says which of the map's it is."""
    from siamang.core.variable import Variable, VariableMap
    from siamang.data import correspondence, drivers, paired, pricing

    rng = np.random.default_rng(4)
    n = 240
    frame = pd.DataFrame(
        {
            "brand": rng.integers(1, 4, n).astype(float),
            "region": rng.integers(1, 4, n).astype(float),
            "y": rng.normal(size=n),
            "a": rng.normal(size=n),
            "b": rng.normal(size=n),
            "b1": rng.integers(0, 2, n).astype(float),
            "b2": rng.integers(0, 2, n).astype(float),
            "b3": rng.integers(0, 2, n).astype(float),
            **{f"gg{p}": rng.integers(0, 2, n).astype(float) for p in (5, 10)},
        }
    )
    variables = VariableMap()
    variables.add_many(
        [
            Variable("brand", "nominal", label="Brand", labels={1: "A", 2: "B", 3: "C"}),
            Variable("region", "nominal", label="Region", labels={1: "N", 2: "S", 3: "W"}),
            Variable("y", "interval", label="Liking"),
        ]
    )
    data = SurveyData(frame=frame, variables=variables)
    mapped = correspondence.analyze(data, "brand", column="region")
    keys = drivers.analyze(data, "y", ["a", "b"])
    cochran = paired.cochran(data, ["b1", "b2", "b3"])
    prices = pricing.gabor_granger(data, ["gg5", "gg10"], prices=[5, 10])
    report = Report(title="Methods").heading("Methods")
    for table in (mapped.table, mapped.rows, mapped.columns, keys.table, cochran.table):
        report.add(table)
    report.add(prices.table).add(prices.curves)
    book = openpyxl.load_workbook(report.save_tables(tmp_path / "m.xlsx"))
    described = [row[2] for row in _rows(book["Contents"])[4:]]
    assert described == [
        "Perceptual map: Brand × Region — inertia",
        "Perceptual map: Brand × Region — rows (Brand)",
        "Perceptual map: Brand × Region — columns (Region)",
        "Key drivers: Liking",
        "Paired tests: Cochran's Q",
        "Price sensitivity: Gabor-Granger — demand and revenue",
        "Price sensitivity: Gabor-Granger — curves",
    ]

"""How a report writes a p-value (ReportTheme.p_values).

"exact" — the default — writes a p as it is kept, byte for byte what reports
wrote before the setting existed (the goldens under tests/documents were
written by the engine before it). "0.01" and "0.001" write one below the
threshold as "< 0.01" / "< 0.001": in every table's cells and footer, in a
report's statistics lines, its Markdown as its HTML, its Excel workbook (a
number format, the cell keeps its number) and its charts' notes ("< .01" in
their APA style). The p-values kept — frames, stats, exports — never change.
"""

from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

from siamang.data import SurveyData
from siamang.reporting import Report, ReportTheme, ReportThemeError, p_values
from siamang.reporting.model_tables import regression_table
from siamang.reporting.result_table import ResultTable

GOLDEN = Path(__file__).resolve().parent / "documents"


@pytest.fixture(autouse=True)
def _no_house_theme(monkeypatch):
    # A report without a theme reads SIAMANG_REPORT_THEME; these tests say
    # which theme they mean.
    monkeypatch.delenv("SIAMANG_REPORT_THEME", raising=False)


def _data() -> SurveyData:
    """120 respondents, built without a random generator so the numbers (and
    the goldens) are the same on every machine: a score that rises with the
    region, a second one that follows it, and a gender that leans by region."""

    i = np.arange(120)
    region = i % 3 + 1
    gender = np.where((i // 3) % 4 == 0, 3 - (region > 1) - 1, (region > 1) + 1)
    wobble = ((i * 37) % 11) / 5.0
    score = region * 0.9 + wobble
    other = score * 0.5 + ((i * 13) % 7) / 3.0
    noise = ((i * 29) % 13) / 4.0
    frame = pd.DataFrame(
        {"region": region, "gender": gender, "score": score, "other": other, "noise": noise}
    )
    return SurveyData(frame)


def _report(theme=None) -> Report:
    """Every kind of table that shows a p-value, a bare frame, and statistics
    lines — no chart, whose pictures differ between matplotlib versions."""

    data = _data()
    report = Report(title="P values", theme=theme)
    report.heading("Tests")
    report.add(data.report.crosstab("gender", "region"), caption="Crosstab")
    report.add(
        data.report.means("other", by="region", method="anova", posthoc="tukey"),
        caption="Group means",
    )
    report.add(data.report.ttest("score", by="gender"), caption="t-test")
    report.add(
        data.report.correlation_matrix(
            ["score", "other", "noise"], method="pearson", adjust="holm", layout="pairs"
        ),
        caption="Correlations",
    )
    model = data.analysis.regression("score", ["other", "noise"])
    report.add(regression_table(model, data), caption="Regression")
    report.add(
        ResultTable(
            data=data,
            frame=pd.DataFrame(
                {"Driver": ["other", "noise"], "Beta": [0.91, 0.02], "Beta p": [2.4e-19, 0.4121]}
            ),
            footer={"Bartlett p": 0.0004, "Fit p": 0.2311},
        ),
        caption="Drivers",
    )
    report.add(
        pd.DataFrame({"term": ["a", "b"], "p_value": [1.2345678e-09, 0.034567891]}),
        caption="Bare",
    )
    report.add({"statistic": 12.5, "p_value": 0.00193, "n": 120}, caption="Kruskal-Wallis")
    report.add({"p": 0.004, "lower": 0.001, "upper": 0.012, "n": 250}, caption="Proportion")
    report.add({"posthoc": "Dunn", "1 vs 3": "z = -3.108, p = 0.0057"}, caption="Pairs")
    return report


# ─── the setting ─────────────────────────────────────────────────────────────


def test_the_setting_is_an_enum_of_the_theme_that_is_exact_by_default():
    from siamang.reporting.theme import _ENUMS

    assert ReportTheme().p_values == "exact"
    assert _ENUMS["p_values"] == ("exact", "0.01", "0.001")
    assert ReportTheme(p_values="0.01").to_dict() == {"p_values": "0.01"}
    assert ReportTheme.from_dict({"p_values": "0.001"}).p_values == "0.001"
    # From a JSON file a threshold may be the number it is; the theme keeps its name.
    assert ReportTheme.from_dict({"p_values": 0.01}) == ReportTheme(p_values="0.01")
    for bad in ("0.05", "< 0.01", 0.05, True, None):
        with pytest.raises(ReportThemeError, match="p_values"):
            ReportTheme(p_values=bad)


def test_the_names_that_hold_p_values():
    named = ["p", "p (Holm)", "p (unadjusted)", "p adjusted", "Beta p", "Bartlett p", "Fit p"]
    named += ["p_value", "p-value", "lr_p", "p_adjusted"]
    assert all(p_values.is_p(name) for name in named)
    for other in ("p adjustment", "p method", "N", "Top", "Group", "Pair", 3, None):
        assert not p_values.is_p(other), other


def test_below_the_threshold_and_not():
    assert p_values.is_below(0.0099, "0.01") and p_values.is_below(0.0, "0.01")
    assert not p_values.is_below(0.01, "0.01")  # p < .01 is false at .01
    assert p_values.is_below(0.0009, "0.001") and not p_values.is_below(0.0009, "exact")
    assert p_values.is_below("< 1e-07", "0.01") and not p_values.is_below("< 0.05", "0.01")
    for value in (float("nan"), None, True, "", "n/a", -0.001):
        assert not p_values.is_below(value, "0.01")
    assert p_values.in_text("z = -3.108, p = 0.0057", "0.01") == "z = -3.108, p < 0.01"
    assert p_values.in_text("z = 1.2, p = 0.2301", "0.01") == "z = 1.2, p = 0.2301"
    assert p_values.in_text("differ at p < 0.05", "0.01") == "differ at p < 0.05"
    assert p_values.excel_format("0.01") == '[<0.01]"< 0.01";General'
    assert p_values.excel_format("exact") is None


def test_a_chart_writes_p_in_its_apa_style_at_the_threshold():
    with p_values.showing("exact"):
        assert [p_values.apa(p) for p in (0.0123, 0.0042, 0.0004)] == [".012", ".004", "< .001"]
    with p_values.showing("0.01"):
        assert [p_values.apa(p) for p in (0.0123, 0.0042, 0.0004)] == [".012", "< .01", "< .01"]
    with p_values.showing("0.001"):
        assert [p_values.apa(p) for p in (0.0123, 0.0042, 0.0004)] == [".012", ".004", "< .001"]


# ─── the renderers ───────────────────────────────────────────────────────────


def test_a_table_outside_a_report_writes_p_as_kept_unless_given_a_theme():
    table = _data().report.crosstab("gender", "region")
    p = table.stats["p"]
    assert p < 0.001
    kept = f"p = {p}"
    assert kept in table.to_markdown() and kept in table.to_html()
    assert table.to_markdown(theme=ReportTheme()) == table.to_markdown()
    assert table.to_html(theme={"p_values": "exact"}) == table.to_html()
    assert "; p < 0.01;" in table.to_markdown(theme={"p_values": "0.01"})
    assert "; p < 0.001;" in table.to_html(theme=ReportTheme(p_values="0.001"))
    # Only the writing: the statistic and the frame keep the p.
    assert table.stats["p"] == p
    assert kept in table.to_markdown()


def test_group_means_and_its_post_hoc_pairs():
    table = _data().report.means("score", by="region", method="anova", posthoc="tukey")
    pairs = table.posthoc_table.to_frame()
    exact = table.to_markdown()
    shown = table.to_markdown(theme={"p_values": "0.01"})
    assert "; p < 0.01;" in shown and "; p = " not in shown.split("**Post-hoc")[0]
    rows = [line for line in shown.splitlines() if line.startswith("| ") and " vs " in line]
    for row, p in zip(rows, pairs["p"], strict=True):
        written = row.strip("|").split("|")[-1].strip()
        assert written == ("< 0.01" if p_values.is_below(p, "0.01") else str(p)), (row, p)
    assert exact == table.to_markdown()
    html = table.to_html(theme={"p_values": "0.001"})
    assert "<td>&lt; 0.001</td>" in html
    assert table.posthoc_table.to_frame().equals(pairs)


def test_tukey_s_floor_becomes_the_report_s_bound(tmp_path):
    """Below what SciPy computes the studentized range to, a Tukey p is the
    text "< 1e-07"; under a threshold it is the threshold's bound like any
    other, and the footer's clause about the floor goes with it."""

    from siamang.data import inference
    from siamang.reporting.stat_tables import PostHocTable

    groups = [np.arange(40.0) % 5, np.arange(40.0) % 5 + 40, np.arange(40.0) % 5 + 80]
    result = inference.posthoc(groups, ["A", "B", "C"], "tukey")
    table = PostHocTable(data=_data(), result=result)
    assert "< 1e-07" in table.to_markdown() and "where it is smaller" in table.to_markdown()
    shown = table.to_markdown(theme={"p_values": "0.01"})
    assert "< 1e-07" not in shown and "| < 0.01 |" in shown
    assert "where it is smaller" not in shown
    assert "where it is smaller" in table.stats["p"]

    # In the report's workbook the floor is text, and so is the bound.
    means = _data().report.means("score", by="region", method="anova", posthoc="tukey")
    report = Report(theme=ReportTheme(p_values="0.01")).add(means, caption="Means")
    sheet = openpyxl.load_workbook(report.save_tables(tmp_path / "t.xlsx"))["Means – Post-hoc"]
    column = [cell.value for cell in sheet[1]].index("p") + 1
    assert [sheet.cell(row=r, column=column).value for r in (2, 3)] == ["< 0.01", "< 0.01"]
    assert sheet.cell(row=4, column=column).number_format == '[<0.01]"< 0.01";General'


def test_t_test_correlations_and_regression():
    data = _data()
    ttest = data.report.ttest("score", by="gender")
    assert f"p = {ttest.stats['p']}" in ttest.to_markdown()
    assert "; p < 0.01;" in ttest.to_markdown(theme={"p_values": "0.01"})

    pairs = data.report.correlation_matrix(
        ["score", "other", "noise"], method="pearson", adjust="holm", layout="pairs"
    )
    frame = pairs.to_frame()
    shown = pairs.to_html(theme={"p_values": "0.01"})
    cells = re.findall(r"<td>([^<]*)</td>", shown)
    width = len(frame.columns)
    for i, row in frame.iterrows():
        for column in ("p", "p (Holm)"):
            written = cells[i * width + list(frame.columns).index(column)]
            assert written == ("&lt; 0.01" if row[column] < 0.01 else str(row[column]))
    # The matrix writes a coefficient and its marks, no p.
    matrix = data.report.correlation_matrix(["score", "other", "noise"], method="pearson")
    assert matrix.to_markdown(theme={"p_values": "0.01"}) == matrix.to_markdown()

    model = data.analysis.regression("score", ["other", "noise"])
    table = regression_table(model, data)
    shown = table.to_markdown(theme={"p_values": "0.001"})
    row = next(line for line in shown.splitlines() if line.startswith("| other"))
    assert row.strip("|").split("|")[4].strip() == "< 0.001"
    assert table.to_frame()["p"].iloc[1] < 0.001  # kept


def test_a_result_table_s_p_columns_and_footer():
    table = ResultTable(
        data=_data(),
        frame=pd.DataFrame({"Driver": ["a", "b"], "Beta p": [2.4e-19, 0.4121]}),
        footer={"Bartlett p": 0.0004, "Fit p": 0.2311, "p-value": "exact"},
    )
    shown = table.to_markdown(theme={"p_values": "0.001"})
    assert "| a | < 0.001 |" in shown and "| b | 0.4121 |" in shown
    assert "Bartlett p < 0.001; Fit p = 0.2311; p-value = exact" in shown


def test_a_table_whose_columns_are_answers_is_left_alone():
    """A crosstab's columns are named by the answers' labels: one called "p"
    holds counts and percentages, and a count of 0 is not "< 0.01"."""

    from siamang.core.variable import Variable, VariableMap

    variables = VariableMap()
    variables.add(Variable(name="answer", scale="nominal", labels={1: "p", 2: "q"}))
    frame = pd.DataFrame({"group": [1, 1, 1, 2, 2, 2], "answer": [1, 1, 1, 2, 2, 2]})
    labeled = SurveyData(frame, variables=variables)
    table = labeled.report.crosstab("group", "answer")
    grid = table.to_markdown().split("\n\n")[0]
    assert grid.splitlines()[0] == "| group | p | q | Total |"
    assert "| 2 | 0 | 3 | 3 |" in grid  # no one in group 2 answered "p"
    assert table.to_markdown(theme={"p_values": "0.01"}).split("\n\n")[0] == grid

    data = SurveyData(_data().frame.assign(region=_data().frame["region"] - 1))
    means = data.report.descriptives(["noise"], by="region", layout="means")
    assert means.to_markdown(theme={"p_values": "0.01"}) == means.to_markdown()


def test_a_report_s_statistics_line_and_its_bare_frames():
    report = _report(ReportTheme(p_values="0.01"))
    markdown = report.to_markdown()
    assert "*Kruskal-Wallis*: statistic = 12.5; p_value < 0.01; n = 120" in markdown
    # A proportion's p is the proportion, not a test's.
    assert "*Proportion*: p = 0.004; lower = 0.001" in markdown
    # A p the engine wrote into a sentence.
    assert "1 vs 3 = z = -3.108, p < 0.01" in markdown
    # A bare frame: the bound, the rest as tabulate writes a number; the
    # HTML writes the cells the Markdown does.
    assert "| a      | < 0.01    |" in markdown and "| b      | 0.0345679 |" in markdown
    html = report.to_html(standalone=True)
    assert (
        "<td>&lt; 0.01</td>\n    </tr>\n    <tr>\n      <td>b</td>\n      <td>0.0345679</td>"
        in (html)
    )
    assert "statistic = 12.5; p_value &lt; 0.01; n = 120" in html


# ─── the whole report, each way ──────────────────────────────────────────────


@pytest.mark.parametrize("mode", ["exact", "0.01", "0.001"])
def test_report_golden(mode):
    """The report's Markdown and HTML in each setting. The exact ones were
    written by the engine before the setting existed, with no theme: the
    default writes what it always wrote, byte for byte."""

    for theme in [None, ReportTheme(), {"p_values": "exact"}] if mode == "exact" else [None]:
        report = _report(theme if mode == "exact" else ReportTheme(p_values=mode))
        markdown, html = report.to_markdown(), report.to_html(standalone=True)
        for text, suffix in ((markdown, "md"), (html, "html")):
            golden = GOLDEN / f"report-p-values.{mode}.{suffix}"
            if not golden.exists():  # pragma: no cover - first run writes the golden
                golden.write_text(text, encoding="utf-8")
            assert text == golden.read_text("utf-8"), f"run: rm {golden} and re-run to accept"


@pytest.mark.parametrize(
    ("mode", "bound", "shown"),
    [
        ("exact", None, None),
        ("0.01", "< 0.01", '[<0.01]"< 0.01";General'),
        ("0.001", "< 0.001", '[<0.001]"< 0.001";General'),
    ],
)
def test_report_end_to_end(tmp_path, mode, bound, shown):
    """Saved as the Save report node saves it: Markdown, HTML and the workbook
    of its tables, one setting for all three."""

    report = _report(ReportTheme(p_values=mode))
    md = report.save(tmp_path / "r.md").read_text("utf-8")
    html = report.save(tmp_path / "r.html").read_text("utf-8")
    book = openpyxl.load_workbook(report.save_tables(tmp_path / "r.xlsx"))
    means = book["Group means"]
    rows = {means.cell(row=r, column=1).value: means.cell(row=r, column=2) for r in range(1, 30)}
    posthoc = book["Group means – Post-hoc"]
    header = [cell.value for cell in posthoc[1]]
    first_p, second_p = (posthoc.cell(row=r, column=header.index("p") + 1) for r in (2, 3))
    drivers = book["Drivers"]
    if bound is None:
        for written in ("p < 0.01;", "p < 0.001;", "| < 0.01 |", "| < 0.001 |"):
            assert written not in md and written not in html
        assert rows["p"].number_format == "General" and first_p.number_format == "General"
        return
    assert f"p {bound}" in md and f"| {bound} |" in md
    assert f"p {bound}" in html and f"<td>{bound.replace('<', '&lt;')}</td>" in html
    # The cells keep their numbers; a number format shows them as the bound.
    assert isinstance(rows["p"].value, float) and rows["p"].value < 0.001
    assert rows["p"].number_format == shown
    assert rows["F"].number_format == "General"
    for cell in (first_p, second_p):  # 0.0053, 2.66e-07
        assert isinstance(cell.value, float) and cell.number_format == shown
    beta_p = drivers.cell(row=2, column=3)
    assert beta_p.value == 2.4e-19 and beta_p.number_format == shown
    assert drivers.cell(row=2, column=2).number_format == "General"  # Beta
    stats = {
        drivers.cell(row=r, column=1).value: drivers.cell(row=r, column=2) for r in range(4, 8)
    }
    assert stats["Bartlett p"].value == 0.0004 and stats["Bartlett p"].number_format == shown
    assert stats["Fit p"].number_format == shown  # General above the threshold


# ─── charts ──────────────────────────────────────────────────────────────────


def test_a_chart_s_note_follows_the_report_that_shows_it(tmp_path):
    """A t-test's Result chart writes its p under the chart when it is drawn —
    at its node, as kept. A report with a threshold draws it again, in its
    notes and in its interactive form; one that writes p as kept does not."""

    from siamang.reporting import chart_theme
    from siamang.reporting.result_charts import chart

    ttest = _data().report.ttest("score", by="gender")
    drawn = chart(ttest)
    drawn._ensure_built()
    p = ttest.stats["p"]
    assert drawn._p_drawn_with == "exact"
    assert drawn._drawn.notes[-1].endswith(f"p = {p}")

    assert chart_theme.in_report(drawn, ReportTheme()) is drawn
    shown = chart_theme.in_report(drawn, ReportTheme(p_values="0.01"))
    assert shown is not drawn and shown._p_drawn_with == "0.01"
    assert shown._drawn.notes[-1].endswith("p < 0.01")
    assert "p < 0.01" in str(shown.vega_lite())
    # Kept for the report's next rendering; the chart itself is as drawn.
    assert chart_theme.in_report(drawn, ReportTheme(p_values="0.01")) is shown
    assert drawn._drawn.notes[-1].endswith(f"p = {p}")

    # A chart that writes no p is never drawn again for it.
    bars = _data().plot.bar("region")
    bars._ensure_built()
    assert bars._p_drawn_with == ""
    assert chart_theme.in_report(bars, ReportTheme(p_values="0.01")) is bars

    report = Report(title="C", theme=ReportTheme(p_values="0.001")).add(drawn, caption="t")
    report.to_markdown(asset_dir=tmp_path)
    assert drawn._redrawn._p_drawn_with == "0.001"
    assert drawn._redrawn._drawn.notes[-1].endswith("p < 0.001")


def test_a_correlation_heatmap_s_tooltips_write_p_at_the_threshold():
    import json

    from siamang.reporting import chart_theme
    from siamang.reporting.result_charts import chart

    def tooltips(drawn):
        return re.findall(r'"detail_0": "([^"]+)"', json.dumps(drawn.vega_lite()))

    data = _data()
    frame = data.frame.assign(close=data.frame["noise"] + ((np.arange(120) * 7) % 9) / 1.5)
    table = SurveyData(frame).report.correlation_matrix(
        ["score", "other", "noise", "close"], method="pearson"
    )
    heatmap = chart(table)
    exact = tooltips(heatmap)
    assert "< .001" in exact and ".011" in exact, exact
    assert tooltips(chart_theme.in_report(heatmap, ReportTheme())) == exact
    shown = tooltips(chart_theme.in_report(heatmap, ReportTheme(p_values="0.01")))
    assert shown == ["< .01" if p == "< .001" or p.startswith(".00") else p for p in exact]


def test_a_flow_s_save_report_sets_it(tmp_path):
    """The Save report node's Look (theme) takes p_values: checked, run, and
    written into the generated script."""

    from siamang.flow import FlowRunner, check_flow, generate_flow
    from siamang.model import from_document, loads

    document = loads((GOLDEN / "brand_awareness.questionnaire.json").read_text("utf-8"))
    survey = from_document(document).survey

    def flow(theme):
        nodes = [
            ("sim", "source.simulated", {"n": 300, "seed": 3}),
            ("xt", "analyze.crosstab", {"row": "region", "col": "gender"}),
            ("means", "analyze.means", {"y": "satisfaction", "by": "region"}),
            ("sec", "output.report_section", {"heading": "Tables"}),
            (
                "save",
                "output.save_report",
                {"title": "T", "path": "outputs/t.md", "xlsx": True, "theme": theme},
            ),
        ]
        edges = [
            ("sim", "data", "xt", "data"),
            ("sim", "data", "means", "data"),
            ("xt", "table", "sec", "items"),
            ("means", "table", "sec", "items"),
            ("sec", "report", "save", "sections"),
        ]
        return {
            "schema_version": "1.0",
            "name": "p",
            "nodes": [{"id": i, "type": t, "params": p} for i, t, p in nodes],
            "edges": [
                {"from": {"node": a, "port": ap}, "to": {"node": b, "port": bp}}
                for a, ap, b, bp in edges
            ],
        }

    [issue] = check_flow(flow({"p_values": "0.05"}), questionnaire=document)
    assert issue.node == "save" and "p_values: '0.05' is not one of exact, 0.01, 0.001." in (
        issue.message
    )
    good = flow({"p_values": "0.01"})
    assert check_flow(good, questionnaire=document) == []
    assert 'theme={"p_values": "0.01"})' in generate_flow(good)

    result = FlowRunner(good, questionnaire=survey, questionnaire_document=document).run(
        cwd=tmp_path
    )
    assert result.ok
    table = result.output("xt", "table")
    written = (tmp_path / "outputs" / "t.md").read_text("utf-8")
    html = (tmp_path / "outputs" / "t.html").read_text("utf-8")
    p = table.stats["p"]
    if p < 0.01:
        assert "; p < 0.01;" in written and "; p < 0.01;" in html
    else:
        assert f"; p = {p};" in written
    book = openpyxl.load_workbook(tmp_path / "outputs" / "t.xlsx")
    formats = {
        cell.number_format
        for sheet in book.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if sheet.cell(row=cell.row, column=1).value == "p" and cell.column == 2
    }
    assert formats == {'[<0.01]"< 0.01";General'}

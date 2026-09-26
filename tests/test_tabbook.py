"""The tab book: every question crossed by the banner, one sheet each.

Its cells are the Banner table's and the crosstab's for the same data — the
same counts, percentages and letters — with the codebook's missing codes left
out; the workbook is opened with openpyxl and read back cell by cell.
"""

from __future__ import annotations

import os
import re
import subprocess
import sys
from datetime import UTC, datetime
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from siamang.core.variable import MissingValue, Variable, VariableMap
from siamang.data.survey_data import SurveyData
from siamang.reporting.tabbook import tabulate, write_tabbook
from siamang.reporting.tables import BannerTable, CrossTable

DOCUMENTS = Path(__file__).resolve().parent / "documents"
ROOT = Path(__file__).resolve().parents[1]
CREATED = datetime(2026, 9, 1, 8, 30, tzinfo=UTC)

REGION = Variable("region", "nominal", label="Region", labels={1: "North", 2: "South", 3: "East"})
GENDER = Variable("gender", "nominal", label="Gender", labels={1: "Man", 2: "Woman"})
SAT = Variable(
    "sat",
    "ordinal",
    label="Satisfaction",
    labels={1: "Poor", 2: "Fair", 3: "Good", 4: "Very good", 5: "Excellent"},
)


def _data(frame: pd.DataFrame, *variables: Variable) -> SurveyData:
    codebook = VariableMap()
    codebook.add_many(list(variables))
    return SurveyData(frame=frame, variables=codebook)


def _complete(seed: int = 7, n: int = 420) -> SurveyData:
    """Everyone answered everything: the Banner table's bases are the tab book's."""
    rng = np.random.default_rng(seed)
    region = rng.choice([1, 2, 3], n, p=[0.45, 0.35, 0.2])
    shift = np.where(region == 1, 0.8, np.where(region == 2, 0.0, -0.4))
    sat = np.clip(np.round(rng.normal(3 + shift, 1.1)), 1, 5).astype(int)
    frame = pd.DataFrame(
        {
            "region": region,
            "gender": rng.choice([1, 2], n),
            "sat": sat,
            "w": rng.choice([0.5, 1.0, 1.5, 2.5], n),
        }
    )
    return _data(frame, REGION, GENDER, SAT, Variable("w", "ratio", label="Weight"))


def _parse(cell: str) -> tuple[float, float, str]:
    """A Banner table cell, ``"45.0% (12) BC"``, as (percent, count, letters)."""
    match = re.fullmatch(r"(-?[\d.]+)% \(([\d.]+)\)(?: ([A-Z]+))?", cell)
    assert match, cell
    return float(match.group(1)), float(match.group(2)), match.group(3) or ""


def _sheet(path: Path, name: str) -> list[tuple]:
    sheet = load_workbook(path)[name]
    return [row for row in sheet.iter_rows(values_only=True)]


def _find(rows: list[tuple], label: str) -> int:
    return next(index for index, row in enumerate(rows) if row and row[0] == label)


# ── the numbers ────────────────────────────────────────────────────────────


@pytest.mark.parametrize("weighted", [False, True])
def test_the_cells_are_the_banner_tables(weighted, tmp_path):
    data = _complete()
    if weighted:
        data = data.with_weight("w")
    banner = BannerTable(data=data, rows=["sat"], columns=["region", "gender"]).to_frame()
    book = tabulate(data, banner=["region", "gender"], questions=["sat"])
    tab = book.tabs[0]
    assert [column.header for column in tab.columns] == [
        "Total",
        "North (A)",
        "South (B)",
        "East (C)",
        "Man (D)",
        "Woman (E)",
    ]
    assert list(banner.columns[2:]) == [
        f"{block}: {column.header}"
        for block, column in zip(["Region"] * 3 + ["Gender"] * 2, tab.columns[1:], strict=True)
    ]
    letters_seen = 0
    for position, (_code, label) in enumerate(tab.answers):
        row = banner[banner["Answer"] == label].iloc[0]
        for at in range(1, len(tab.columns)):
            percent, count, letters = _parse(row.iloc[1 + at])
            assert round(tab.column_percent[position][at] * 100, 1) == percent
            assert round(tab.counts[position][at], 1) == pytest.approx(count, abs=0.051)
            assert tab.letters[position][at] == letters
            letters_seen += bool(letters)
    assert letters_seen >= 3  # the data has differences to find
    base = banner[banner["Question"] == "Base"].iloc[0]
    assert [round(value, 1) for value in (tab.weighted_base or tab.base)[1:]] == [
        pytest.approx(float(value)) for value in base.iloc[2:]
    ]

    path = write_tabbook(
        data,
        tmp_path / "book.xlsx",
        banner=["region", "gender"],
        questions=["sat"],
        created=CREATED,
    )["Workbook"]
    rows = _sheet(Path(path), "sat")
    header = rows[5]
    assert header[1] is None and header[2] == "North (A)" and header[4] == "South (B)"
    at = _find(rows, "Very good")
    position = [label for _code, label in tab.answers].index("Very good")
    # Counts on the answer's line, the column percentage under it, its letters beside it.
    assert rows[at][2] == pytest.approx(tab.counts[position][1])
    assert rows[at + 1][2] == pytest.approx(tab.column_percent[position][1])
    assert rows[at + 1][3] == (tab.letters[position][1] or None)
    base_rows = [row for row in rows if row[0] in ("Base", "Base (unweighted)", "Base (weighted)")]
    assert base_rows[0][1] == len(data.frame)
    assert len(base_rows) == (2 if weighted else 1)


@pytest.mark.parametrize("weighted", [False, True])
def test_the_column_percentages_are_the_crosstabs(weighted):
    data = _complete(seed=3)
    if weighted:
        data = data.with_weight("w")
    cross = CrossTable(data=data, row="sat", col="region", pct="col", test=False).to_frame()
    counts = CrossTable(data=data, row="sat", col="region", pct="none", test=False).to_frame()
    tab = tabulate(data, banner=["region"], questions=["sat"]).tabs[0]
    for position, (_code, label) in enumerate(tab.answers):
        row = cross[cross["Satisfaction"] == label].iloc[0]
        assert [round(tab.column_percent[position][at] * 100, 1) for at in (1, 2, 3)] == [
            row["North"],
            row["South"],
            row["East"],
        ]
        # The crosstab's Total column is the answer's count over every column.
        total = counts[counts["Satisfaction"] == label].iloc[0]["Total"]
        assert tab.counts[position][0] == pytest.approx(float(total), abs=0.051)


def test_letters_by_hand():
    """Left 60 of 100 say yes, Right 40 of 100: z = 0.2 / √(0.5·0.5·(2/100)) =
    2.83, p = 0.0047 — Left's yes is higher than Right's (B), Right's no than
    Left's (A)."""
    rows = [{"grp": 1, "ans": 1}] * 60 + [{"grp": 1, "ans": 2}] * 40
    rows += [{"grp": 2, "ans": 1}] * 40 + [{"grp": 2, "ans": 2}] * 60
    grp = Variable("grp", "nominal", label="Group", labels={1: "Left", 2: "Right"})
    ans = Variable("ans", "nominal", label="Answer", labels={1: "Yes", 2: "No"})
    data = _data(pd.DataFrame(rows), grp, ans)
    tab = tabulate(data, banner=["grp"], questions=["ans"]).tabs[0]
    assert tab.letters == [["", "B", ""], ["", "", "A"]]
    assert tabulate(data, banner=["grp"], questions=["ans"], level=0.001).tabs[0].letters == [
        ["", "", ""],
        ["", "", ""],
    ]
    assert tabulate(data, banner=["grp"], questions=["ans"], letters=False).tabs[0].letters == [
        ["", "", ""],
        ["", "", ""],
    ]
    # Below the minimum base a column is not tested.
    small = tabulate(data, banner=["grp"], questions=["ans"], min_base=101).tabs[0]
    assert small.letters == [["", "", ""], ["", "", ""]]


def test_missing_codes_are_left_out_and_said(tmp_path):
    sat = Variable(
        "sat",
        "ordinal",
        label="Satisfaction",
        labels={1: "Poor", 2: "Good", 9: "Don't know"},
        missing=(MissingValue(9, "Don't know", kind="dont_know"),),
    )
    region = Variable(
        "region",
        "nominal",
        label="Region",
        labels={1: "North", 2: "South", 99: "Refused"},
        missing=(MissingValue(99, "Refused", kind="refusal"),),
    )
    frame = pd.DataFrame(
        {
            "region": [1, 1, 1, 2, 2, 99, 99, 1],
            "sat": [1, 2, 9, 2, 2, 1, 9, None],
        }
    )
    data = _data(frame, region, sat)
    book = tabulate(data, banner=["region"], questions=["sat"])
    tab = book.tabs[0]
    # Out of the base: the two "Don't know" and the blank. The two "Refused"
    # regions are in Total and in no column.
    assert tab.answers == [(1, "Poor"), (2, "Good")]
    assert [column.header for column in tab.columns] == ["Total", "North (A)", "South (B)"]
    assert tab.base == [5, 2, 2]
    assert tab.counts == [[2.0, 1.0, 0.0], [3.0, 1.0, 2.0]]
    assert tab.column_percent[1] == [pytest.approx(0.6), 0.5, 1.0]
    assert tab.left_out == [(9, 2)]
    path = tmp_path / "missing.xlsx"
    stat = write_tabbook(data, path, banner=["region"], questions=["sat"], created=CREATED)
    rows = _sheet(path, "sat")
    notes = [row[0] for row in rows if row[0] and str(row[0]).startswith("Left out")]
    assert notes == ["Left out, as missing codes: Satisfaction: 2 (9 = Don't know)."]
    notes = dict((row[0], row[1]) for row in _sheet(path, "Notes") if row[0])
    assert notes["Missing codes left out"] == (
        "Satisfaction: 2 (9 = Don't know)\n"
        "Region: 2 (99 = Refused) — in Total, in no column of the banner"
    )
    assert stat["Sheets written"] == 1 and stat["Questions skipped"] == 0


def test_row_percentages_add_up_across_each_banner_variable(tmp_path):
    data = _complete(seed=11, n=200)
    book = tabulate(data, banner=["region", "gender"], questions=["sat"], percentages="row")
    tab = book.tabs[0]
    for row in tab.row_percent:
        assert row[0] == 1.0
        assert sum(row[1:4]) == pytest.approx(1.0) and sum(row[4:6]) == pytest.approx(1.0)
    path = tmp_path / "row.xlsx"
    stat = write_tabbook(
        data, path, banner=["region"], questions=["sat"], percentages="row", counts=False
    )
    rows = _sheet(path, "sat")
    at = _find(rows, "Poor")
    # Counts off: one line per answer, its row percentages — and no letters,
    # which compare column percentages: no column is kept for them, and no
    # header names its column by a letter.
    assert rows[at][1] == 1.0 and rows[at + 1][0] == "Fair"
    assert rows[5][1:5] == (None, "North", "South", "East")
    sheet = load_workbook(path)["sat"]
    assert sheet.cell(row=at + 1, column=3).number_format == "0.0%"
    no_letters = (
        "no letters: they compare column percentages, and this book shows percentages of the row"
    )
    assert stat["Test"] == no_letters
    assert f"{no_letters.capitalize()}." in [row[0] for row in rows]
    notes = {row[0]: row[1] for row in _sheet(path, "Notes") if row[0]}
    assert notes["Test"] == no_letters and "Alpha" not in notes


def test_a_multiple_choice_question(tmp_path):
    aware = Variable(
        "aware", "nominal", label="Brands known", labels={1: "Acme", 2: "Globex", 3: "Initech"}
    )
    grp = Variable("grp", "nominal", label="Group", labels={1: "Left", 2: "Right"})
    lists = [[1, 2]] * 20 + [[1]] * 20 + [[3]] * 10 + [[]] * 5
    lists += [[2]] * 30 + [[2, 3]] * 10 + [[1, 3]] * 5
    groups = [1] * 55 + [2] * 45
    data = _data(pd.DataFrame({"grp": groups, "aware": lists}), grp, aware)
    tab = tabulate(data, banner=["grp"], questions=["aware"]).tabs[0]
    assert tab.multiple
    # The five empty answers are no answer: 50 answered on the Left.
    assert tab.base == [95, 50, 45]
    assert tab.counts == [[45.0, 40.0, 5.0], [60.0, 20.0, 40.0], [25.0, 10.0, 15.0]]
    assert sum(row[0] for row in tab.column_percent) > 1  # several answers each
    # Acme: 80 % against 11 % — Left is higher; Globex: 40 % against 89 %.
    assert tab.letters[0] == ["", "B", ""] and tab.letters[1] == ["", "", "A"]
    path = tmp_path / "multi.xlsx"
    write_tabbook(data, path, banner=["grp"], questions=["aware"], created=CREATED)
    rows = _sheet(path, "aware")
    assert "multiple answers" in rows[1][0]
    assert any(
        row[0] == "Several answers were allowed, so the percentages add to more than 100%."
        for row in rows
    )


def test_means_of_an_interval_question_weighted_by_hand():
    score = Variable("score", "interval", label="Score")
    grp = Variable("grp", "nominal", label="Group", labels={1: "Left", 2: "Right"})
    frame = pd.DataFrame(
        {"grp": [1, 1, 1, 2, 2], "score": [2.0, 4.0, 5.0, 1.0, 3.0], "w": [1, 1, 2, 1, 3]}
    )
    data = _data(frame, grp, score, Variable("w", "ratio"))
    tab = tabulate(data, banner=["grp"], questions=["score"]).tabs[0]
    # Five values, no labels: counted by value, and the mean beside them.
    assert [label for _code, label in tab.answers] == ["1", "2", "3", "4", "5"]
    assert tab.means["Mean"][1] == pytest.approx(11 / 3)
    assert tab.means["Standard deviation"][1] == pytest.approx(np.std([2, 4, 5], ddof=1))
    weighted = tabulate(data.with_weight("w"), banner=["grp"], questions=["score"]).tabs[0]
    # Left: (2 + 4 + 10) / 4 = 4; variance (4 + 0 + 2) / 4 · 3/2 = 2.25.
    assert weighted.means["Mean"][1] == pytest.approx(4.0)
    assert weighted.means["Standard deviation"][1] == pytest.approx(1.5)
    # Right: (1 + 9) / 4 = 2.5; variance (2.25 + 3 · 0.25) / 4 · 2/1 = 1.5.
    assert weighted.means["Mean"][2] == pytest.approx(2.5)
    assert weighted.means["Standard deviation"][2] == pytest.approx(1.5**0.5)
    without = tabulate(data, banner=["grp"], questions=["score"], means=False).tabs[0]
    assert without.means is None


def test_the_default_questions_and_what_is_left_out():
    rng = np.random.default_rng(1)
    n = 60
    frame = pd.DataFrame(
        {
            "respondent_id": [f"r{i}" for i in range(n)],
            "region": rng.choice([1, 2], n),
            "sat": rng.choice([1, 2, 3, 4, 5], n),
            "age": rng.integers(18, 80, n),
            "comment": [f"text {i}" for i in range(n)],
            "aware": [[1, 2]] * n,
            "w": 1.0,
            "url_source": "panel",
        }
    )
    data = _data(
        frame,
        Variable("respondent_id", "nominal", role="id"),
        Variable("region", "nominal", label="Region", labels={1: "North", 2: "South"}),
        SAT,
        Variable("age", "ratio", label="Age"),
        Variable("comment", "nominal", label="Comment"),
        Variable("aware", "nominal", label="Aware", labels={1: "A", 2: "B"}),
        Variable("w", "ratio", role="weight"),
        Variable("url_source", "nominal"),
        Variable("q_gone", "nominal", label="Not asked"),
    ).with_weight("w")
    book = tabulate(data, banner=["region"])
    assert [tab.variable for tab in book.tabs] == ["sat", "aware"]
    assert book.skipped == [
        ("q_gone", "not in the data"),
        (
            "comment",
            "60 different answers and no answer labels — an open answer? Code it first "
            "(Code open answers) or band it (Bands)",
        ),
    ]
    # Age is ratio: in the book only when named, then with its mean.
    named = tabulate(data, banner=["region"], questions=["age", "nope"])
    assert named.tabs[0].answers == [] and named.tabs[0].means is not None
    assert named.skipped == [("nope", "not in the data")]
    off = tabulate(data, banner=["region"], questions=["age"], means=False)
    assert off.skipped[0][1].startswith("ratio with ")
    assert off.skipped[0][1].endswith("no answers to count — turn on Means")


def test_sheet_names_are_unique_short_and_legal(tmp_path):
    long = "a_very_long_variable_name_of_forty_chars"
    names = [long, long.upper(), "q[1]:x", "notes", "History"]
    frame = pd.DataFrame({name: [1, 2, 1, 2] for name in names} | {"g": [1, 1, 2, 2]})
    data = _data(
        frame,
        *(Variable(name, "nominal", label=f"Question {name}") for name in names),
        Variable("g", "nominal", label="G", labels={1: "One", 2: "Two"}),
    )
    path = tmp_path / "names.xlsx"
    write_tabbook(data, path, banner=["g"], questions=names, created=CREATED)
    workbook = load_workbook(path)
    assert workbook.sheetnames == [
        "Contents",
        long[:31],
        long.upper()[:29] + "~2",
        "q_1__x",
        "notes~2",
        "History~2",
        "Notes",
    ]
    contents = workbook["Contents"]
    links = [
        (cell.value, cell.hyperlink.location)
        for row in contents.iter_rows(min_row=5, max_row=9, min_col=2, max_col=2)
        for cell in row
    ]
    assert links[1] == (f"Question {long.upper()}", f"'{long.upper()[:29]}~2'!A1")


def test_the_workbook_contents_notes_and_formats(tmp_path):
    data = _complete(seed=5, n=300).with_weight("w")
    path = tmp_path / "book" / "tabs.xlsx"  # the folder is made
    stat = write_tabbook(
        data,
        path,
        banner=["region", "gender"],
        questions=["sat"],
        correction="bonferroni",
        created=CREATED,
    )
    assert stat == {
        "Workbook": str(path),
        "Sheets written": 1,
        "Questions skipped": 0,
        "Banner": "Region, Gender",
        "Percentages": "column",
        "Test": "two-sided z-test of column proportions at 0.05, within each banner variable, "
        "Bonferroni-corrected",
        "Weight": "w",
    }
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Contents", "sat", "Notes"]
    contents = workbook["Contents"]
    assert contents["A1"].value == "Tab book"
    assert contents["B5"].value == "Satisfaction"
    assert contents["B5"].hyperlink.location == "'sat'!A1"
    assert contents["D5"].value == 300
    sheet = workbook["sat"]
    assert sheet["A1"].value == "Satisfaction"
    assert sheet["A3"].hyperlink.location == "'Contents'!A1"
    assert sheet.freeze_panes == "B9"  # below the two header rows and the two bases
    assert sheet["B5"].value == "Total" and sheet["C5"].value == "Region"
    assert "C5:H5" in {str(merged) for merged in sheet.merged_cells.ranges}
    assert sheet["A7"].value == "Base (unweighted)" and sheet["A8"].value == "Base (weighted)"
    assert sheet["B9"].number_format == "#,##0" and sheet["B10"].number_format == "0.0%"
    assert sheet.column_dimensions["D"].width == 5  # a letters column
    notes = {row[0]: row[1] for row in workbook["Notes"].iter_rows(values_only=True) if row[0]}
    assert notes["Created"] == "2026-09-01 08:30 UTC"
    assert notes["Weight"] == "w" and notes["Respondents"] == 300
    assert notes["Banner"] == "Total, Region (A–C), Gender (D–E)"
    assert notes["Alpha"] == 0.05
    assert notes["Bonferroni"].startswith("yes")
    assert notes["Minimum base for a test"] == "30 respondents (Kish's effective base)"
    assert notes["Missing codes left out"] == "none met"


def test_what_a_tab_book_refuses(tmp_path):
    data = _complete(n=40)
    with pytest.raises(ValueError, match="must end in .xlsx"):
        write_tabbook(data, tmp_path / "book.csv", banner=["region"])
    with pytest.raises(ValueError, match="nothing to show"):
        tabulate(data, banner=["region"], percentages="none", counts=False)
    with pytest.raises(ValueError, match="banner variable 'nope' is not in the data"):
        tabulate(data, banner=["nope"])
    listed = data.with_frame(data.frame.assign(region=[[1]] * 40))
    with pytest.raises(ValueError, match="Region holds multiple-choice answers"):
        tabulate(listed, banner=["region"], questions=["sat"])
    bare = SurveyData(frame=data.frame)
    with pytest.raises(ValueError, match="no codebook"):
        tabulate(bare, banner=["region"])


# ── the node ───────────────────────────────────────────────────────────────


@pytest.fixture(scope="module")
def questionnaire_doc():
    from siamang.model import loads

    return loads((DOCUMENTS / "brand_awareness.questionnaire.json").read_text("utf-8"))


@pytest.fixture(scope="module")
def survey(questionnaire_doc):
    from siamang.model import from_document

    return from_document(questionnaire_doc).survey


def _tabbook_flow(params: dict | None = None) -> dict:
    nodes = [
        ("src", "source.responses", {}),
        (
            "cell",
            "prepare.cell_weights",
            {"variable": "gender", "targets": {"1": 0.5, "2": 0.4, "3": 0.1}},
        ),
        ("apply", "prepare.apply_weight", {}),
        ("book", "output.tabbook", {"banner": ["region", "gender"], **(params or {})}),
    ]
    edges = [
        ("src", "data", "cell", "data"),
        ("cell", "data", "apply", "data"),
        ("apply", "data", "book", "data"),
    ]
    return {
        "schema_version": "1.0",
        "name": "tabs",
        "nodes": [{"id": i, "type": t, "params": p} for i, t, p in nodes],
        "edges": [
            {"from": {"node": a, "port": ap}, "to": {"node": b, "port": bp}}
            for a, ap, b, bp in edges
        ],
    }


def test_the_node_is_registered_and_checked(questionnaire_doc):
    from siamang.flow import check_flow, default_registry

    spec = default_registry().get("output.tabbook")
    assert spec.title == "Tab book (Excel)"
    assert spec.outputs == {"stat": "Stat"}
    assert spec.params["percentages"].values == ("column", "row", "none")
    assert spec.params["path"].default == "outputs/tabbook.xlsx"
    # Level and Multiple comparisons are read only with the letters.
    assert not spec.reads("level", {"letters": False})
    assert spec.reads("level", {"letters": True})
    assert check_flow(_tabbook_flow(), questionnaire=questionnaire_doc) == []
    issues = check_flow(
        _tabbook_flow({"percentages": "none", "counts": False, "letters": False}),
        questionnaire=questionnaire_doc,
    )
    assert [(issue.severity, issue.code) for issue in issues] == [("error", "PARAM_CONFLICT")]
    issues = check_flow(_tabbook_flow({"percentages": "row"}), questionnaire=questionnaire_doc)
    assert [(issue.severity, issue.message) for issue in issues] == [
        (
            "warning",
            "book: Significance letters compare column percentages, so the sheets show them "
            "only with Percentages column; with row percentages or counts only they are left out.",
        )
    ]
    issues = check_flow(_tabbook_flow({"banner": ["age"]}), questionnaire=questionnaire_doc)
    assert [issue.code for issue in issues] == ["VARIABLE_SCALE"]
    help_text = default_registry().get("prepare.apply_weight").params["column"].help
    weighted, _unweighted = help_text.split("Unweighted, and saying so:")
    assert "Trend" in weighted and "Tab book (Excel)" in weighted


def _cells(path: Path) -> dict[str, list[tuple]]:
    workbook = load_workbook(path)
    return {
        name: [
            row
            for row in workbook[name].iter_rows(values_only=True)
            if not (name == "Notes" and row[0] == "Created")
        ]
        for name in workbook.sheetnames
    }


def test_the_node_runs_and_its_script_writes_the_same_workbook(questionnaire_doc, survey, tmp_path):
    from siamang.codegen import generate_questionnaire
    from siamang.flow import FlowRunner, generate_flow
    from siamang.io import write_snapshot

    responses = survey.simulate(n=320, seed=4)
    flow = _tabbook_flow({"questions": ["satisfaction", "aware", "trust_acme", "age"]})
    runner_dir, script_dir = tmp_path / "runner", tmp_path / "script"
    result = FlowRunner(flow, questionnaire=survey, questionnaire_document=questionnaire_doc).run(
        sources={"src": responses}, cwd=runner_dir
    )
    assert result.ok
    stat = result.output("book")
    assert stat["Sheets written"] == 4 and stat["Weight"] == "weight"
    workbook = runner_dir / "outputs" / "tabbook.xlsx"
    ours = _cells(workbook)
    assert list(ours) == ["Contents", "satisfaction", "aware", "trust_acme", "age", "Notes"]
    trust = ours["trust_acme"]
    assert trust[0][0] == "Trust: Acme"
    assert any(
        row[0] and str(row[0]).startswith("Left out, as missing codes: Trust: Acme")
        for row in trust
    )
    assert "Refused" not in [row[0] for row in trust]
    assert [row[0] for row in ours["age"] if row[0] in ("Mean", "Standard deviation")] == [
        "Mean",
        "Standard deviation",
    ]

    code = generate_flow(flow, questionnaire_doc)
    assert "from siamang.reporting.tabbook import write_tabbook" in code
    assert "letters=True" in code and 'correction="none"' in code
    no_letters = generate_flow(_tabbook_flow({"letters": False}), questionnaire_doc)
    assert "letters=False" in no_letters and "level=" not in no_letters
    (script_dir / "survey").mkdir(parents=True)
    (script_dir / "survey" / "__init__.py").write_text("", encoding="utf-8")
    (script_dir / "survey" / "questionnaire.py").write_text(
        generate_questionnaire(questionnaire_doc), encoding="utf-8"
    )
    (script_dir / "scripts").mkdir()
    script = script_dir / "scripts" / "tabs.py"
    script.write_text(code, encoding="utf-8")
    snapshot = write_snapshot(responses, script_dir / "data" / "responses.csv")
    env = {**os.environ, "PYTHONPATH": os.pathsep.join([str(script_dir), str(ROOT)])}
    completed = subprocess.run(
        [sys.executable, str(script), "--data", str(snapshot)],
        cwd=script_dir,
        env=env,
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr
    assert _cells(script_dir / "outputs" / "tabbook.xlsx") == ours

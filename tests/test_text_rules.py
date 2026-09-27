"""Codeframe version 2: a coder's decisions, and rules for the answers nobody decided.

Every case here is written out by hand: an answer, the rule, and what a person
reading both would expect. The rules are small on purpose — words, ``*`` for
word forms, ``|`` for alternatives, ``not_`` for a negated mention, ``~N`` for
words near each other — so that what they do can be predicted, and these tests
are that prediction.
"""

from __future__ import annotations

import json
import time

import pandas as pd
import pytest

from siamang.core import Variable, VariableMap
from siamang.data import SurveyData, text_coding, text_rules

fp = text_coding.fingerprint


def _cf(*themes, **extra) -> dict:
    payload = {"schema_version": "2.0", "variable": "why", "themes": list(themes)}
    payload.update(extra)
    return payload


def _theme(code, label, include=(), **more) -> dict:
    rules = {k: more.pop(k) for k in ("require", "exclude", "scope") if k in more}
    entry = {"code": code, "label": label, **more}
    if include or rules:
        entry["rules"] = {"include": list(include), **rules}
    return entry


def _codes(text, payload) -> list[int]:
    return text_coding.explain(text, payload)["codes"]


def _matches(term: str, text: str, **theme) -> bool:
    """Whether one theme whose only include term is ``term`` codes ``text``."""
    return _codes(text, _cf(_theme(1, "T", [term], **theme))) == [1]


# ─── words and terms ─────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("term", "text", "expected"),
    [
        # a word, whatever its case and the punctuation around it
        ("late", "It was LATE!", True),
        ("late", "latest news", False),
        # * for the letters of a word's forms, anywhere in it
        ("delay*", "Delayed twice", True),
        ("delay*", "delays", True),
        ("delay*", "delay", True),
        ("delay*", "a relay", False),
        ("*ing", "charging", True),
        ("*arg*", "recharge", True),
        ("c*ng", "charging", True),
        ("c*ng", "charge", False),
        ("a*b*c", "abxxc", True),
        ("a*b*c", "acb", False),
        # | for alternatives at one place
        ("slow|late", "so slow", True),
        ("slow|late", "late again", True),
        ("slow|late", "fast", False),
        # a phrase: its words in order, next to each other
        ("customer service", "The customer service was bad", True),
        ("customer service", "service to the customer", False),
        ("customer service", "customer, service", False),
        ("customer serv*|support", "customer support", True),
        # ~N: within N words of each other, in either order
        ("staff ~2 rude", "the staff were very rude", True),
        ("staff ~1 rude", "the staff were very rude", False),
        ("staff ~0 rude", "rude staff", True),
        ("staff ~3 rude", "rude, and the staff", False),  # not across punctuation
        ("customer service ~2 slow|bad", "slow or bad customer service", True),
        # an apostrophe inside a word is part of it, however it was typed
        ("don't", "I don’t", True),
        ("can't|cannot", "I cannot say", True),
    ],
)
def test_term_forms(term, text, expected):
    assert _matches(term, text) is expected


@pytest.mark.parametrize(
    ("term", "text", "expected"),
    [
        # a term matches a mention that is not negated…
        ("late", "Delivery wasn't late", False),
        ("late", "It was never late", False),
        ("late", "No late deliveries", False),
        ("late", "not really late", False),
        ("late", "Delivery was not late", False),
        # …and not_ one that is
        ("not_late", "Delivery wasn't late", True),
        ("not_late", "It was never late", True),
        ("not_late", "It was late", False),
        # a negation spelled in the term is the term's own
        ("not late", "Delivery was not late", True),
        ("don't know", "I don't know", True),
        ("don't know", "I don't really know", False),
        ("don't ~2 know", "I don't really know", True),
        ("no problems", "No problems at all", True),
        ("problem*", "No problems at all", False),
        # every n't form, typed with or without its apostrophe
        ("helpful", "The staff isn't helpful", False),
        ("helpful", "The staff didnt seem helpful", False),
        ("helpful", "They can't be helpful", False),
        ("helpful", "They cannot be helpful", False),
        ("helpful", "without helpful staff", False),
        ("helpful", "hardly helpful", False),
        # three words reach…
        ("late", "not at all late", False),
        ("late", "not at all very late", True),
        # …stopped by punctuation, and, or, yet, but and the other clause words
        ("fast", "not cheap, fast", True),
        ("fast", "not cheap and fast", True),
        ("fast", "not cheap or fast", True),
        ("fast", "not cheap yet fast", True),
        ("fast", "not cheap but fast", True),
        ("fast", "not cheap however fast", True),
        ("fast", "not cheap fast", False),
    ],
)
def test_negation(term, text, expected):
    assert _matches(term, text) is expected


def test_a_form_with_many_stars_takes_no_time_on_a_long_word():
    """A pattern language would backtrack here; a form is read in one pass."""
    started = time.perf_counter()
    assert not _matches("a*a*a*a*a*a*a*a*a*a*b", "a" * 20_000)
    assert time.perf_counter() - started < 2


def test_a_negated_mention_is_found_by_not_in_either_word_of_a_proximity():
    assert _matches("staff ~3 not_rude", "the staff were not rude")
    assert not _matches("staff ~3 rude", "the staff were not rude")


@pytest.mark.parametrize(
    ("term", "text", "expected"),
    [
        # The staff are under the negation the term asks for with not_: the
        # term's own, as a negation written in the term is.
        ("not_friendly staff", "No friendly staff", True),
        ("not_friendly staff", "Not friendly staff at all", True),
        ("not_friendly ~3 staff", "Not friendly staff at all", True),
        ("staff ~3 not_friendly", "Not friendly staff", True),  # either order
        ("staff ~3 not_friendly", "The staff were not friendly", True),
        ("not_friendly ~3 staff", "The staff were not friendly", True),
        # …but a word negated by another negation is still negated
        ("not_friendly staff", "Friendly staff", False),
        ("not_late ~3 staff", "not late and no staff", False),
        ("not_late ~5 staff", "never late and no staff", False),
    ],
)
def test_a_not_form_takes_the_words_under_the_negation_it_asks_for(term, text, expected):
    assert _matches(term, text) is expected


@pytest.mark.parametrize(
    ("term", "text", "expected"),
    [
        # An n't form is one however it is typed: with its apostrophe,
        # without it, or spelled out with not.
        ("don't know", "I dont know", True),
        ("don't know", "I do not know", True),
        ("don't know", "I don’t know", True),
        ("do not know", "I don't know", True),
        ("dont know", "do not know", True),
        ("would not recommend", "I wouldn't recommend it", True),
        ("wouldn't recommend", "I would not recommend it", True),
        ("can't|cannot find", "I can not find it", True),
        ("won't", "they will not", True),
        ("don't|doesn't work", "it does not work", True),
        ("didn't|never arrive", "it did not arrive", True),
        ("didn't|never arrive", "it never arrive", True),
        # "not" in a term is every negation written with not
        ("not happy", "I wasn't happy", True),
        ("not happy", "I'm not happy", True),
        ("not happy", "never happy", False),
        # a mention under such a negation is negated, as it was
        ("know", "I do not know", False),
        ("recommend", "would not recommend", False),
        ("not_recommend", "would not recommend", True),
        # "not" after a word that is no auxiliary stays a word of its own
        ("not late", "definitely not late", True),
        ("do not", "do, not", False),
    ],
)
def test_the_nt_forms_are_one_however_they_are_typed(term, text, expected):
    assert _matches(term, text) is expected


def test_do_not_is_read_as_one_negating_word():
    got = text_coding.explain("I do not know, can not say", _cf(_theme(1, "DK", ["don't know"])))
    assert [t["word"] for t in got["tokens"]] == ["i", "don't", "know", "can't", "say"]
    assert [t["negated_by"] for t in got["tokens"]] == [None, None, "don't", None, "can't"]
    assert got["codes"] == [1]
    # In a term whose place holds another word too, "do not" is warned of.
    found = text_coding.validate(_cf(_theme(1, "A", ["do|really not know"])))
    assert any("'do not' is read as 'don't'" in w.message for w in found.warnings)


# ─── clauses and scope ───────────────────────────────────────────────────────


def test_a_rule_reads_one_clause_unless_its_scope_is_the_answer():
    text = "Delivery was quick but the box was damaged"
    clause = _cf(_theme(1, "Damaged delivery", ["damaged"], require=["delivery"]))
    assert _codes(text, clause) == []
    whole = _cf(_theme(1, "Damaged delivery", ["damaged"], require=["delivery"], scope="answer"))
    assert _codes(text, whole) == [1]
    # the codeframe's scope is every theme's unless it says otherwise
    assert _codes(text, {**clause, "scope": "answer"}) == [1]
    explained = text_coding.explain(text, clause)
    assert explained["clauses"] == ["delivery was quick", "the box was damaged"]


@pytest.mark.parametrize(
    "text",
    [
        "Fast. Delivery damaged",
        "fast; delivery damaged",
        "fast (delivery damaged)",
        "fast — delivery damaged",
        "fast, although delivery damaged",
        "fast whereas delivery damaged",
        "fast except delivery damaged",
        "fast plus delivery damaged",
        "fast though delivery damaged",
    ],
)
def test_clauses_end_at_punctuation_and_at_the_clause_words(text):
    explained = text_coding.explain(text, _cf(_theme(1, "T", ["x"])))
    assert explained["clauses"] == ["fast", "delivery damaged"]


@pytest.mark.parametrize(
    "text",
    [
        "fast\ndelivery damaged",
        "fast\r\n\r\n  delivery damaged",
        "fast delivery damaged",
        "fast - delivery damaged",
        "- fast\n- delivery damaged",
        "fast -- delivery damaged",
        "• fast • delivery damaged",
        "fast | delivery damaged",
        "fast / delivery damaged",
        "fast/ delivery damaged",
        "fast · delivery damaged",
    ],
)
def test_a_line_break_a_spaced_dash_a_bullet_and_a_slash_end_a_clause(text):
    explained = text_coding.explain(text, _cf(_theme(1, "T", ["x"])))
    assert explained["clauses"] == ["fast", "delivery damaged"]


def test_a_negation_and_a_clause_rule_stop_at_the_end_of_a_line():
    payload = _cf(
        _theme(1, "Price", ["price"]),
        _theme(2, "Damaged", ["damaged"], require=["staff"]),
        multiple=True,
    )
    assert _codes("Not happy\nPrice too high", payload) == [1]
    assert _codes("Not good - price too high", payload) == [1]
    assert _codes("- not helpful\n- price high", payload) == [1]
    assert _codes("Staff fine\nparcel damaged", payload) == []
    assert _codes("Staff fine, parcel damaged", payload) == []
    assert _codes("staff: parcel damaged", payload) == []
    assert _codes("the staff had the parcel damaged", payload) == [2]
    # A line's words stay together: a phrase does not reach into the next.
    assert not _matches("customer service", "customer\nservice")


@pytest.mark.parametrize(
    ("term", "text"),
    [
        ("e mail", "e-mail"),
        ("n a", "n/a"),
        ("and or", "and/or"),
        ("col legi", "col·legi"),
        ("well known", "well‐known"),
    ],
)
def test_a_joiner_between_two_letters_ends_no_clause(term, text):
    assert _matches(term, text)


def test_an_answer_read_across_lines_keeps_its_fingerprint():
    """A line break ends a clause for the rules, but the fingerprint — and so
    a coder's decision — is the answer's as it always was."""
    text = "Not happy\nPrice too high"
    assert fp(text) == fp("not happy price too high")
    payload = _cf(_theme(1, "Price", ["price"]), assignments={fp("x\ny"): [1]})
    explained = text_coding.explain(text, payload)
    assert explained["fingerprint"] == fp(text)
    assert explained["normalised"] == "not happy\nprice too high"
    assert _codes("x  \n  Y", payload) == [1] and _codes("X y", payload) == [1]
    # The same words on one line and on two are one fingerprint, read apart.
    got = text_coding.preview({text: 2, "Not happy Price too high": 1}, payload)
    assert [(a["fingerprint"], a["codes"], a["count"]) for a in got["answers"]] == [
        (fp(text), [1], 2),
        (fp(text), [], 1),
    ]
    series = pd.Series([text, "Not happy Price too high", "not happy\r\nprice too high"])
    cf = text_coding.parse(payload)
    assert text_coding.codes(series, cf).tolist() == [1, pd.NA, 1]
    assert text_coding.coverage(pd.DataFrame({"why": series}), cf)["by_rules"] == 2


def test_an_exclude_vetoes_only_where_it_is():
    price = _cf(_theme(1, "Price too high", ["price"], exclude=["fair", "reasonable"]))
    assert _codes("The price is fair", price) == []
    assert _codes("The price is high, the service fair", price) == [1]
    assert _codes("price reasonable", price) == []


def test_every_require_group_must_match():
    theme = _theme(1, "Rude driver", ["rude"], require=[["driver", "courier"], ["delivery*"]])
    payload = _cf(theme)
    assert _codes("the courier was rude at delivery", payload) == [1]
    assert _codes("the courier was rude", payload) == []
    assert _codes("rude at delivery", payload) == []
    # A flat list is one group: any of its terms.
    assert _codes("rude driver", _cf(_theme(1, "R", ["rude"], require=["driver", "x"]))) == [1]


# ─── replacements ────────────────────────────────────────────────────────────


def test_replacements_are_whole_words_or_phrases_made_before_the_rules_read():
    payload = _cf(
        _theme(1, "Service", ["customer service"]),
        _theme(2, "Delivery", ["delivery"]),
        _theme(3, "Apps", ["application"]),
        _theme(4, "Helpful", ["helpful"]),
        replace=[
            {"from": "customer care", "to": "customer service"},
            {"from": "delievery", "to": "delivery"},
            {"from": "app", "to": "application"},
            {"from": "dont", "to": "don't"},
            {"from": "N/A", "to": "nothing"},
        ],
    )
    assert _codes("Customer Care was great", payload) == [1]
    assert _codes("the delievery", payload) == [2]
    assert _codes("the app", payload) == [3]
    assert _codes("an apple", payload) == []  # a word, not the start of one
    assert _codes("they dont seem helpful", payload) == []  # now a negation
    assert text_coding.explain("n/a", payload)["normalised"] == "nothing"


# ─── several themes, exclusive themes, priority ──────────────────────────────


THEMES = [
    _theme(1, "Speed", ["slow|fast"], priority=1),
    _theme(2, "Price", ["price*|expensive"], priority=5),
    _theme(3, "Staff", ["staff"], priority=3),
    _theme(4, "Nothing", ["nothing", "no comment", "don't know"], exclusive=True),
    _theme(5, "Other", ["other"], priority=3),
]


def test_several_themes_come_in_the_codeframe_order():
    many = _cf(*THEMES, multiple=True)
    assert _codes("slow, expensive and rude staff", many) == [1, 2, 3]


def test_max_codes_keeps_the_highest_priority_and_ties_go_by_order():
    capped = _cf(*THEMES, multiple=True, max_codes=2)
    assert _codes("slow, expensive and rude staff", capped) == [2, 3]
    # Staff and Other share a priority: the one listed first wins the cut.
    assert _codes("staff and other things, slow", capped) == [3, 5]
    dropped = text_coding.explain("slow, expensive and rude staff", capped)["dropped"]
    assert dropped == [
        {"code": 1, "label": "Speed", "reason": "max_codes keeps 2, and they ranked higher"}
    ]


def test_one_theme_an_answer_keeps_the_best():
    single = _cf(*THEMES)
    assert _codes("slow, expensive and rude staff", single) == [2]
    assert _codes("slow staff", single) == [3]
    assert _codes("staff and other", single) == [3]


def test_an_exclusive_theme_stands_alone_or_not_at_all():
    many = _cf(*THEMES, multiple=True)
    assert _codes("Nothing really", many) == [4]
    assert _codes("Don't know", many) == [4]
    assert _codes("Nothing, except it was slow", many) == [1]
    reason = text_coding.explain("Nothing, except it was slow", many)["dropped"]
    assert reason == [
        {"code": 4, "label": "Nothing", "reason": "exclusive, and another theme matched"}
    ]
    # Two exclusive themes: the one ranked higher.
    two = _cf(
        _theme(1, "Nothing", ["nothing"], exclusive=True),
        _theme(2, "Don't know", ["know"], exclusive=True, priority=2),
    )
    assert _codes("nothing, I know", {**two, "multiple": True}) == [2]


# ─── a coder's decisions ─────────────────────────────────────────────────────


def test_a_coder_decides_before_the_rules_and_may_decide_there_is_no_theme():
    payload = _cf(
        *THEMES,
        multiple=True,
        assignments={
            fp("Slow and expensive"): [3],  # the coder read it differently
            fp("Staff were slow"): [],  # read, and no theme
            fp("Blah"): [1, 5],
        },
    )
    assert text_coding.explain("slow and EXPENSIVE", payload)["codes"] == [3]
    assert text_coding.explain("slow and EXPENSIVE", payload)["source"] == "hand"
    assert text_coding.explain("  staff were slow ", payload)["manual"] == []
    assert text_coding.explain("Staff were slow", payload)["codes"] == []
    assert text_coding.explain("Blah", payload)["codes"] == [1, 5]
    assert text_coding.explain("slow staff", payload)["source"] == "rule"
    assert text_coding.explain("unrelated", payload)["source"] == "uncoded"
    # One theme an answer: a decision naming two keeps the one ranked higher.
    single = {**payload, "multiple": False}
    assert text_coding.explain("Blah", single)["codes"] == [5]


def test_an_old_assignment_keeps_matching_its_answer():
    """The fingerprint is the version 1 one, unchanged."""
    assert fp("  Charging  IS too slow ") == fp("charging is too slow")
    payload = _cf(_theme(1, "Charging"), assignments={fp("Charging is too slow"): 1})
    assert _codes("charging is TOO slow", payload) == [1]


# ─── Unicode ─────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("term", "text", "expected"),
    [
        ("цена", "Цена слишком высокая", True),
        ("цен*", "Высокие цены", True),
        ("café", "Le Café était froid", True),
        ("café", "cafe", False),  # diacritics are kept
        ("straße", "STRASSE", True),  # case-folded as Unicode folds it
        ("नमस्ते", "नमस्ते दुनिया", True),  # a word with combining vowel signs is one word
        ("مرحبا", "مرحبا بكم", True),
        ("价格太高", "价格太高", True),
        ("l'eau", "L’eau", True),
    ],
)
def test_words_in_any_script(term, text, expected):
    assert _matches(term, text) is expected


def test_an_answer_in_another_language_is_answered_and_uncoded_not_lost():
    payload = _cf(_theme(1, "Price", ["price"]))
    result = text_coding.preview({"Цена высокая": 3, "价格太高": 2, "price": 1}, payload)
    assert result["coverage"]["answered"] == 6 and result["coverage"]["uncoded"] == 5
    assert [a["source"] for a in result["answers"]] == ["uncoded", "uncoded", "rule"]
    frame = pd.DataFrame({"why": ["Цена высокая", "价格太高", "price", ""]})
    cf = text_coding.parse(payload)
    assert list(text_coding.uncoded_answers(frame, cf)) == ["Цена высокая", "价格太高"]
    words = text_coding.explain("Цена не высокая", payload)["tokens"]
    assert [w["word"] for w in words] == ["цена", "не", "высокая"]
    assert not any(w["negated"] for w in words)  # the negations are English


# ─── the file ────────────────────────────────────────────────────────────────


FULL = _cf(
    _theme(1, "Late", ["late", "delay*"], group="Delivery", exclude=["not_late"]),
    _theme(2, "Damaged", ["damag*", "broken|broke"], group="Delivery", priority=2),
    _theme(3, "Rude staff", ["rude"], require=[["staff", "driver"]], definition="Manners."),
    _theme(9, "Nothing", ["nothing"], exclusive=True, scope="answer"),
    multiple=True,
    max_codes=3,
    replace=[{"from": "dont", "to": "don't"}],
    assignments={fp("Great!"): [], fp("ok"): 2, fp("both"): [1, 2]},
    sentiment={fp("ok"): 1},
    model="hand",
    built_at="2026-09-27",
    source_rows=10,
)


def test_a_version_2_codeframe_round_trips_through_its_file(tmp_path):
    path = tmp_path / "why.codeframe.json"
    path.write_text(json.dumps(FULL), encoding="utf-8")
    cf = text_coding.load(path)
    assert cf.version == 2 and cf.multiple and cf.max_codes == 3 and cf.scope == "clause"
    assert cf.assignments[fp("Great!")] == () and cf.assignments[fp("both")] == (1, 2)
    assert cf.themes[1].group == "Delivery" and cf.themes[3].exclusive
    assert cf.themes[2].rules.require == (("staff", "driver"),)
    assert cf.nets == {"Delivery": (1, 2)}
    written = cf.to_dict()
    assert written["schema_version"] == "2.0"
    assert written["assignments"][fp("ok")] == 2 and written["assignments"][fp("Great!")] == []
    assert text_coding.parse(written) == cf
    assert text_coding.parse(json.loads(json.dumps(written))) == cf
    # Only fingerprints of answers are kept: no answer's text is in the file.
    text = json.dumps(written, ensure_ascii=False)
    assert "Great!" not in text and '"ok"' not in text
    assert text_coding.parse({**FULL, "schema_version": 2}).version == 2


def test_an_exclusive_theme_without_rules_is_a_valid_theme():
    payload = _cf(_theme(1, "Price", ["price"]), _theme(2, "Nothing", exclusive=True))
    assert text_coding.validate(payload).to_dict() == {"ok": True, "errors": [], "warnings": []}


@pytest.mark.parametrize(
    ("change", "message", "path"),
    [
        ({"themes": []}, "the codeframe has no themes", ["themes"]),
        ({"variable": "drop table"}, "is not a variable name", ["variable"]),
        ({"language": "ru"}, "language 'ru' is not supported", ["language"]),
        ({"scope": "sentence"}, "scope is 'clause' or 'answer'", ["scope"]),
        ({"max_codes": -1}, "max_codes is a whole number", ["max_codes"]),
        ({"multiple": "yes"}, "multiple is true or false", ["multiple"]),
        (
            {"themes": [_theme(1, "A"), _theme(1, "B")]},
            "theme code 1 appears twice",
            ["themes", 1, "code"],
        ),
        ({"themes": [_theme(1, " ")]}, "theme 1 has no label", ["themes", 0, "label"]),
        (
            {"themes": [_theme(2**63, "Big", ["late"])]},
            "theme code 9223372036854775808 is out of range",
            ["themes", 0, "code"],
        ),
        (
            {"themes": [_theme(1, "A"), _theme(-(2**31) - 1, "B")]},
            "a code is a whole number from -2147483648 to 2147483647",
            ["themes", 1, "code"],
        ),
        ({"themes": [_theme(1e300, "Big")]}, "is out of range", ["themes", 0, "code"]),
        (
            {"assignments": {fp("x"): 7}},
            "names unknown theme 7",
            ["assignments", fp("x")],
        ),
        (
            {"assignments": {"The price": 1}},
            "is not an answer's fingerprint",
            ["assignments", "The price"],
        ),
        (
            {"themes": [{"code": 1, "label": "A", "examples": ["The price"]}]},
            "keeps no answers' texts",
            ["themes", 0, "examples"],
        ),
        (
            {"themes": [_theme(1, "A", ["  "])]},
            "the term is empty",
            ["themes", 0, "rules", "include", 0],
        ),
        (
            {"themes": [_theme(1, "A", ["re:late|slow"])]},
            "regular expressions are not supported",
            ["themes", 0, "rules", "include", 0],
        ),
        (
            {"themes": [_theme(1, "A", ["staff ~ rude"])]},
            "~ takes the number of words",
            ["themes", 0, "rules", "include", 0],
        ),
        (
            {"themes": [_theme(1, "A", ["a ~2 b ~3 c"])]},
            "one ~N at most",
            ["themes", 0, "rules", "include", 0],
        ),
        (
            {"themes": [_theme(1, "A", ["~2 rude"])]},
            "needs words on both sides",
            ["themes", 0, "rules", "include", 0],
        ),
        (
            {"themes": [_theme(1, "A", ["*"])]},
            "needs some of its own",
            ["themes", 0, "rules", "include", 0],
        ),
        (
            {"themes": [_theme(1, "A", ["not_"])]},
            "not_ needs a word",
            ["themes", 0, "rules", "include", 0],
        ),
        (
            {"themes": [_theme(1, "A", ["slow||late"])]},
            "empty alternative",
            ["themes", 0, "rules", "include", 0],
        ),
        (
            {"themes": [_theme(1, "A", ["x"], require=["a", ["b"]])]},
            "not both",
            ["themes", 0, "rules", "require"],
        ),
        (
            {"replace": [{"from": "a", "to": "b"}, {"from": " A ", "to": "c"}]},
            "is replaced twice",
            ["replace", 1, "from"],
        ),
        ({"replace": [{"from": "", "to": "b"}]}, "needs the words", ["replace", 0, "from"]),
        ({"schema_version": "3.0"}, "newer than this siamang reads", ["schema_version"]),
    ],
)
def test_an_unusable_codeframe_is_refused_with_a_reason_and_a_place(change, message, path):
    payload = {**_cf(_theme(1, "Price", ["price"])), **change}
    found = text_coding.validate(payload)
    assert not found.ok
    assert any(message in e.message and list(e.path) == path for e in found.errors), found
    with pytest.raises(text_coding.CodeframeError, match="codeframe: "):
        text_coding.parse(payload)


@pytest.mark.parametrize(
    ("theme", "message"),
    [
        (_theme(1, "A", ["e-mail"]), "'-' is not part of a word"),
        (_theme(1, "A", ["'cause"]), "neither begins nor ends with an apostrophe"),
        (_theme(1, "A", ["but"]), "'but' ends a clause"),
        (_theme(1, "A", ["colour"]), "'colour' is replaced by 'color'"),
        (_theme(1, "A", ["not_and"]), "'and' is never negated itself"),
        (_theme(1, "A", ["late", "late"]), "is given twice"),
        (_theme(1, "A", ["late"], exclude=["LATE"]), "both included and excluded"),
        (_theme(1, "A", require=["staff"]), "have no include term"),
        (_theme(1, "A", ["x"], group="Solo"), "the net 'Solo' has one theme"),
        (_theme(1, "A", ["x"], colour="red"), "'colour' is not a theme field"),
    ],
)
def test_a_term_that_can_never_match_is_a_warning(theme, message):
    payload = _cf(theme, replace=[{"from": "colour", "to": "color"}])
    found = text_coding.validate(payload)
    assert found.ok, found.errors
    assert any(message in w.message for w in found.warnings), found.warnings
    text_coding.parse(payload)  # it applies; that term does nothing


def test_a_warning_names_its_theme_when_an_earlier_one_is_refused():
    found = text_coding.validate(
        _cf(
            {"code": "x", "label": "Bad"},
            _theme(1, "A"),
            _theme(1, "A again"),
            _theme(2, "B", group="Solo"),
        )
    )
    net = next(w for w in found.warnings if "the net 'Solo'" in w.message)
    assert list(net.path) == ["themes", 3, "group"]


def test_the_largest_codes_a_variable_holds_apply():
    for code in (2**31 - 1, -(2**31)):
        cf = text_coding.parse(_cf(_theme(code, "Edge", ["late"])))
        assert text_coding.codes(pd.Series(["late"]), cf).tolist() == [code]
    # A version 1 codeframe refuses a code its theme variable cannot hold,
    # rather than failing when it is applied.
    for code in (2**63, float("inf")):
        with pytest.raises(text_coding.CodeframeError, match="codeframe: "):
            text_coding.parse(
                {
                    "schema_version": "1.0",
                    "variable": "why",
                    "themes": [{"code": code, "label": "B"}],
                }
            )
    big = {"schema_version": "1.0", "variable": "why", "themes": [{"code": 2**40, "label": "B"}]}
    assert text_coding.parse(big).themes[0].code == 2**40


def test_version_2_fields_in_a_version_1_file_are_said_to_do_nothing():
    old = {
        "schema_version": "1.0",
        "variable": "why",
        "multiple": True,
        "themes": [{"code": 1, "label": "A", "rules": {"include": ["late"]}}],
    }
    found = text_coding.validate(old)
    assert found.ok and [list(w.path) for w in found.warnings] == [
        ["multiple"],
        ["themes", 0, "rules"],
    ]
    assert "ignored unless schema_version is 2.0" in found.warnings[0].message
    assert text_coding.parse(old).version == 1  # read as version 1 reads it


def test_but_is_a_word_a_rule_over_the_whole_answer_can_match():
    payload = _cf(_theme(1, "A", ["cheap but good"], scope="answer"))
    assert text_coding.validate(payload).warnings == ()
    assert _codes("Cheap but good!", payload) == [1]


def test_a_single_theme_codeframe_warns_of_decisions_with_several():
    found = text_coding.validate(
        _cf(_theme(1, "A"), _theme(2, "B"), assignments={fp("x"): [1, 2]}, max_codes=2)
    )
    messages = " ".join(w.message for w in found.warnings)
    assert "names 2 themes" in messages and "max_codes does nothing" in messages


# ─── applying ────────────────────────────────────────────────────────────────


def _data(answers) -> SurveyData:
    variables = VariableMap()
    variables.add(Variable("why", "nominal", label="Why?", dtype="str"))
    return SurveyData(frame=pd.DataFrame({"why": answers}), variables=variables)


ANSWERS = [
    "Delivery was late and the box damaged",  # rules: 1, 2
    "It was late",  # rules: 1
    "The driver was rude",  # rules: 3
    "Nothing",  # rules: 9
    "ok",  # by hand: 2
    "Great!",  # by hand: no theme
    "Something else entirely",  # uncoded
    "",  # not answered
    None,  # not answered
    "wasn't late",  # uncoded: the only mention is negated
]


def test_a_multiple_codeframe_makes_a_multiple_choice_variable():
    cf = text_coding.parse(FULL)
    out = text_coding.apply(_data(ANSWERS), cf)
    assert out.frame["why_theme"].tolist() == [
        [1, 2],
        [1],
        [3],
        [9],
        [2],
        [],
        None,
        None,
        None,
        None,
    ]
    assert out.variables["why_theme"].labels == {
        1: "Late",
        2: "Damaged",
        3: "Rude staff",
        9: "Nothing",
    }
    assert out.variables["why_theme"].scale == "nominal"
    assert text_coding.sources(pd.Series(ANSWERS), cf).tolist()[:7] == [
        "rule",
        "rule",
        "rule",
        "rule",
        "hand",
        "hand",
        "uncoded",
    ]
    from siamang.data import multi

    assert multi.is_multi(out.frame["why_theme"])


def test_a_single_codeframe_makes_a_nominal_variable():
    cf = text_coding.parse({**FULL, "multiple": False, "max_codes": 0})
    out = text_coding.apply(_data(ANSWERS), cf)
    # Damaged (priority 2) outranks Late; "Great!" has no theme; the rest NA.
    assert out.frame["why_theme"].tolist()[:6] == [2, 1, 3, 9, 2, pd.NA]
    assert str(out.frame["why_theme"].dtype) == "Int64"


def test_coverage_counts_what_was_decided_by_hand_and_by_the_rules():
    cf = text_coding.parse(FULL)
    assert text_coding.coverage(_data(ANSWERS).frame, cf) == {
        "answered": 8,
        "coded": 6,
        "uncoded": 2,
        "by_hand": 2,
        "by_rules": 4,
        "no_theme": 1,
    }
    assert list(text_coding.uncoded_answers(_data(ANSWERS).frame, cf)) == [
        "Something else entirely",
        "wasn't late",
    ]


def test_a_missing_answer_is_blank_however_the_column_holds_it():
    """pd.NA and NaT are not the texts "<NA>" and "NaT": nobody answered."""
    payload = _cf(
        _theme(1, "Late", ["late"]),
        _theme(9, "Nothing / N/A", ["nothing", "na", "n a", "nat"], exclusive=True),
    )
    cf = text_coding.parse(payload)
    strings = pd.Series(["late", None, "n/a", pd.NA], dtype="string")
    assert text_coding.codes(strings, cf).tolist() == [1, pd.NA, 9, pd.NA]
    assert text_coding.sources(strings, cf).tolist() == ["rule", pd.NA, "rule", pd.NA]
    assert text_coding.coverage(pd.DataFrame({"why": strings}), cf)["answered"] == 2
    mixed = pd.Series(["late", pd.NaT, float("nan"), None, "nothing"], dtype="object")
    assert text_coding.coverage(pd.DataFrame({"why": mixed}), cf)["answered"] == 2
    table = _data(strings.astype(object)).report.themes(cf).to_frame().set_index("Theme")
    assert table.loc["Nothing / N/A", "N"] == 1 and table.loc["Nothing / N/A", "%"] == 50.0
    assert text_coding.preview(strings, cf)["coverage"]["answered"] == 2
    assert text_coding.explain(pd.NA, cf)["text"] == ""
    assert text_coding.normalise(pd.NA) == "" and text_coding.normalise(pd.NaT) == ""
    # A version 1 codeframe no longer counts them as answered either.
    old = text_coding.parse(
        {"schema_version": "1.0", "variable": "why", "themes": [{"code": 1, "label": "L"}]}
    )
    assert text_coding.coverage(pd.DataFrame({"why": strings}), old)["answered"] == 2
    assert list(text_coding.uncoded_answers(pd.DataFrame({"why": strings}), old)) == [
        "late",
        "n/a",
    ]


def test_a_coders_several_themes_come_in_the_codeframe_order():
    payload = _cf(*THEMES, multiple=True, assignments={fp("Beta"): [3, 1], fp("Gamma"): [5, 2]})
    cf = text_coding.parse(payload)
    assert text_coding.codes(pd.Series(["Beta", "gamma"]), cf).tolist() == [[1, 3], [2, 5]]
    assert text_coding.explain("beta", cf)["codes"] == [1, 3]
    # The file keeps them as the coder gave them.
    assert cf.to_dict()["assignments"][fp("Beta")] == [3, 1]
    # One theme an answer: the one ranked highest (priority, then order).
    single = text_coding.parse({**payload, "multiple": False})
    assert text_coding.codes(pd.Series(["Beta", "gamma"]), single).tolist() == [3, 2]


def test_answers_collected_after_the_rules_are_coded_by_them():
    """The rules run at every run: an answer nobody has read gets its theme."""
    cf = text_coding.parse(FULL)
    later = ANSWERS + ["The parcel arrived broken", "delays, delays"]
    assert text_coding.codes(pd.Series(later), cf).tolist()[-2:] == [[2], [1]]


# ─── the table ───────────────────────────────────────────────────────────────


def test_the_theme_table_counts_respondents_with_nets_and_how_they_were_coded():
    cf = text_coding.parse(FULL)
    table = text_coding.apply(_data(ANSWERS), cf).report.themes(cf)
    frame = table.to_frame()
    # Largest first, ties by label; a net's themes under it.
    assert frame["Theme"].tolist() == [
        "Delivery (net)",
        "Damaged",
        "Late",
        "Nothing",
        "Rude staff",
        "No theme",
        "Coded",
        "Coded by hand",
        "Coded by rules",
        "Uncoded",
    ]
    # The net counts the respondent with both of its themes once: 3, not 4.
    assert frame["N"].tolist() == [3, 2, 2, 1, 1, 1, 6, 2, 4, 2]
    assert frame["%"].tolist() == [37.5, 25.0, 25.0, 12.5, 12.5, 12.5, 75.0, 25.0, 50.0, 25.0]
    stats = table.stats
    assert stats["Answered"] == 8 and stats["Coded by hand"] == 2 and stats["Coded by rules"] == 4
    assert stats["Coverage"] == "75.0 % of the answers are coded"
    assert stats["Distinct uncoded answers"] == 2
    assert stats["Percentages"] == (
        "of the respondents who answered; a respondent can have several themes, so the "
        "themes add up to more than 100 %"
    )
    assert stats["Nets"] == "a net counts a respondent once, however many of its themes they have"
    assert table.theme_rows()["Theme"].tolist() == ["Damaged", "Late", "Nothing", "Rude staff"]
    single = text_coding.parse({**FULL, "multiple": False})
    stats = _data(ANSWERS).report.themes(single).stats
    assert stats["Percentages"] == "of the respondents who answered"
    assert "Weight" not in stats


def test_the_theme_table_on_weighted_data_says_it_counts_people():
    cf = text_coding.parse(FULL)
    data = _data(ANSWERS)
    weighted = data.with_frame(data.frame.assign(w=[2.0] * len(ANSWERS))).with_weight("w")
    table = weighted.report.themes(cf)
    assert table.to_frame()["N"].tolist()[0] == 3
    assert table.stats["Weight"] == "unweighted (the weight 'w' is not applied)"


def test_sentiment_splits_themes_and_nets_of_a_version_2_codeframe():
    cf = text_coding.parse(FULL)
    table = _data(ANSWERS).report.themes(cf, sentiment=True)
    rows = table.to_frame().set_index("Theme")
    assert rows.loc["Damaged", "Positive %"] == 100.0  # "ok", the one scored
    assert pd.isna(rows.loc["Late", "Positive %"])
    assert table.stats["Sentiment"] == (
        "negative 0.0 %, neutral 0.0 %, positive 100.0 % of 1 answer"
    )
    plain = text_coding.parse({**FULL, "sentiment": {}})
    assert _data(ANSWERS).report.themes(plain, sentiment=True).stats["Sentiment"] == (
        "not in this codeframe"
    )


def test_a_theme_label_goes_to_excel_as_text_never_as_a_formula(tmp_path):
    from openpyxl import load_workbook

    payload = _cf(
        _theme(1, "=SUM(1,2)", ["price"], group="=A1"), _theme(2, "B", ["b"], group="=A1")
    )
    cf = text_coding.parse(payload)
    path = _data(["price", "b"]).report.themes(cf).export_xlsx(tmp_path / "themes.xlsx")
    cells = [c.value for row in load_workbook(path).active.iter_rows() for c in row]
    assert "=SUM(1,2)" in cells and "=A1 (net)" in cells
    kinds = [c.data_type for row in load_workbook(path).active.iter_rows() for c in row]
    assert "f" not in kinds


def test_a_chart_of_the_table_draws_the_themes():
    from siamang.reporting import result_charts as rc

    cf = text_coding.parse(FULL)
    chart = rc.chart(_data(ANSWERS).report.themes(cf))
    assert chart.png()[:4] == b"\x89PNG"


# ─── what an editor calls ────────────────────────────────────────────────────


def test_preview_codes_each_distinct_answer_and_counts_themes_nets_and_coverage():
    answers = {
        "Delivery was late and the box damaged": 3,
        "delivery was LATE and the box damaged": 1,  # the same answer
        "The driver was rude, the box broken": 2,
        "ok": 4,
        "Great!": 1,
        "?": 1,
        "": 9,
    }
    result = text_coding.preview(answers, FULL)
    assert [a["count"] for a in result["answers"]] == [4, 2, 4, 1, 1]
    first = result["answers"][0]
    assert first["source"] == "rule" and first["codes"] == [1, 2]
    assert first["fingerprint"] == fp("Delivery was late and the box damaged")
    assert first["hits"] == [
        {
            "code": 1,
            "label": "Late",
            "term": "late",
            "fragment": "late",
            "clause": "delivery was late and the box damaged",
        },
        {
            "code": 2,
            "label": "Damaged",
            "term": "damag*",
            "fragment": "damaged",
            "clause": "delivery was late and the box damaged",
        },
    ]
    second = result["answers"][1]
    assert second["codes"] == [2, 3] and second["hits"][1]["clause"] == "the driver was rude"
    assert [a["source"] for a in result["answers"][2:]] == ["hand", "hand", "uncoded"]
    themes = {t["code"]: t for t in result["themes"]}
    assert themes[2] == {
        "code": 2,
        "label": "Damaged",
        "group": "Delivery",
        "count": 10,
        "by_hand": 4,
        "by_rules": 6,
        "negated": 0,
        "percent": 83.3,
    }
    assert result["nets"] == [{"group": "Delivery", "codes": [1, 2], "count": 10, "percent": 83.3}]
    assert result["coverage"] == {
        "answered": 12,
        "coded": 11,
        "uncoded": 1,
        "by_hand": 5,
        "by_rules": 6,
        "no_theme": 1,
        "percent_coded": 91.7,
    }
    assert result["distinct"]["answered"] == 5
    assert json.dumps(result)  # JSON as it is
    # Pairs, a Series and a parsed codeframe say the same.
    assert text_coding.preview(list(answers.items()), text_coding.parse(FULL)) == result
    series = pd.Series([text for text, n in answers.items() for _ in range(n)])
    assert text_coding.preview(series, FULL)["coverage"] == result["coverage"]


def test_explain_says_why_step_by_step():
    got = text_coding.explain("Delivery wasn't late, but the driver was RUDE", FULL)
    assert got["normalised"] == "delivery wasn't late, but the driver was rude"
    assert got["clauses"] == ["delivery wasn't late", "the driver was rude"]
    late = next(t for t in got["tokens"] if t["word"] == "late")
    assert late == {"word": "late", "negated": True, "negated_by": "wasn't", "clause": 0}
    assert got["manual"] is None
    by_theme = {(r["code"], r["status"]) for r in got["rules"]}
    assert (1, "negation") in by_theme and (3, "fired") in by_theme
    negation = next(r for r in got["rules"] if r["status"] == "negation")
    assert negation["reason"].startswith("the words are negated here")
    assert got["codes"] == [3] and got["source"] == "rule"
    vetoed = text_coding.explain("rude people", FULL)["rules"]
    assert vetoed == [
        {
            "code": 3,
            "label": "Rude staff",
            "scope": "clause",
            "clause": "rude people",
            "status": "vetoed",
            "term": "rude",
            "fragment": "rude",
            "reason": "it requires one of 'staff', 'driver', and none is in this clause",
        }
    ]
    excluded = text_coding.explain(
        "late, but not late", _cf(_theme(1, "Late", ["late"], exclude=["not_late"], scope="answer"))
    )["rules"][0]
    assert excluded["status"] == "vetoed" and "excludes 'not_late'" in excluded["reason"]
    manual = text_coding.explain("OK", FULL)
    assert manual["manual"] == [2] and manual["source"] == "hand" and manual["dropped"] == []


def test_preview_counts_the_answers_a_theme_loses_to_a_negation():
    """A negated mention does not match — the owner's rule — and preview says
    what that costs each theme, so a coder can read those answers."""
    payload = _cf(
        _theme(1, "Parcel", ["parcel"]),
        _theme(2, "Staff", ["staff"]),
        _theme(3, "Delivery", ["deliver*"], scope="answer"),
        _theme(4, "Late", ["late"]),
        multiple=True,
        assignments={fp("no staff, no parcel"): [2]},
    )
    answers = {
        "never received my parcel": 3,
        "Not enough staff, and the parcel was fine": 2,
        "no information about delivery": 1,
        "late, but the parcel wasn't late": 1,
        "no staff, no parcel": 5,  # a coder decided: the rules do not run
        "late": 1,
    }
    got = text_coding.preview(answers, payload)
    negated = {t["code"]: t["negated"] for t in got["themes"]}
    assert negated == {1: 3, 2: 2, 3: 1, 4: 0}
    rows = {a["text"]: a for a in got["answers"]}
    assert rows["never received my parcel"]["source"] == "uncoded"
    assert rows["never received my parcel"]["negated"] == [
        {"code": 1, "label": "Parcel", "term": "parcel", "fragment": "parcel"}
    ]
    assert rows["Not enough staff, and the parcel was fine"]["codes"] == [1]
    assert [n["code"] for n in rows["Not enough staff, and the parcel was fine"]["negated"]] == [2]
    # A theme it got elsewhere in the answer is not lost to the negation.
    assert rows["late, but the parcel wasn't late"]["codes"] == [1, 4]
    assert rows["late, but the parcel wasn't late"]["negated"] == []
    assert rows["no staff, no parcel"]["negated"] == []
    # coding() carries the same, for a table or an editor to use.
    coded = text_coding.coding(["never received my parcel"], text_coding.parse(payload))[0]
    assert [(theme.code, term, words) for theme, term, words in coded.negated] == [
        (1, "parcel", "parcel")
    ]


def test_suggest_offers_frequent_words_and_phrases_of_the_uncoded_answers():
    answers = {
        "The app crashes all the time": 3,
        "app crashes": 2,
        "wasn't fast": 2,
        "not fast at all, but fine": 1,
        "late": 5,  # the codeframe codes it: left out
    }
    got = text_coding.suggest(answers, 3, codeframe=FULL)
    assert got["words"] == [
        {"term": "app", "count": 5, "example": "The app crashes all the time"},
        {"term": "crashes", "count": 5, "example": "The app crashes all the time"},
        {"term": "not_fast", "count": 3, "example": "wasn't fast"},
    ]
    assert got["phrases"][0] == {
        "term": "app crashes",
        "count": 5,
        "example": "The app crashes all the time",
    }
    everything = text_coding.suggest(answers, 30)
    assert {"term": "late", "count": 5, "example": "late"} in everything["words"]
    terms = {w["term"] for w in everything["words"]} | {p["term"] for p in everything["phrases"]}
    assert not terms & {"the", "all", "but", "time the", "wasn't"}
    assert text_coding.suggest(answers, 30, min_count=6) == {"words": [], "phrases": []}
    # A negated mention is counted apart: a word can come both ways.
    both = text_coding.suggest({"not fast": 3, "fast": 1, "never fast enough": 2, "so fast": 2})
    assert [(w["term"], w["count"]) for w in both["words"]][:2] == [("not_fast", 5), ("fast", 3)]
    # "do not know" is suggested as it is matched: don't know.
    assert text_coding.suggest({"I do not know": 2, "dont know": 1})["phrases"] == [
        {"term": "don't know", "count": 3, "example": "I do not know"}
    ]


def test_preview_is_fast_on_a_large_study():
    """50,000 distinct answers, 30 themes of 10 terms each, in a few seconds
    on one CPU: the rules are compiled once and each word looked up once."""
    import random

    rng = random.Random(7)
    vocab = [f"w{i}" for i in range(3000)]
    fill = "the a was is it and but not never very really so i my".split()
    themes = []
    for code in range(1, 31):
        w = rng.sample(vocab, 11)
        themes.append(
            _theme(
                code,
                f"T{code}",
                [
                    w[0] + "*",
                    f"{w[1]}|{w[2]}",
                    f"{w[3]} {w[4]}",
                    f"{w[5]} ~3 {w[6]}",
                    f"not_{w[7]}",
                    w[8],
                    "*" + w[9][1:],
                    w[10],
                    f"q{code}",
                    f"z{code} y{code}",
                ],
                require=[rng.sample(vocab, 3)] if code % 5 == 0 else None,
                group=f"G{code % 6}",
            )
        )
    for theme in themes:
        if theme["rules"].get("require") is None:
            theme["rules"].pop("require", None)
    payload = _cf(*themes, multiple=True, max_codes=3)
    answers: dict[str, int] = {}
    while len(answers) < 50_000:
        words = [
            rng.choice(vocab) if rng.random() < 0.5 else rng.choice(fill)
            for _ in range(rng.randint(3, 20))
        ]
        answers[" ".join(words) + rng.choice([".", "!", ""])] = rng.randint(1, 3)
    started = time.perf_counter()
    result = text_coding.preview(answers, payload)
    elapsed = time.perf_counter() - started
    assert len(result["answers"]) == 50_000 and result["coverage"]["by_rules"] > 0
    assert elapsed < 15, f"{elapsed:.1f} s"


def test_a_long_word_costs_no_more_than_a_short_one():
    """An answer with no spaces — a pasted string, a keyboard mash — is one
    word; looking it up in the forms with * does not grow with its square."""
    payload = _cf(_theme(1, "T", ["delay*", "*ing", "c*ng", "a*b*c", "*arg*"]))
    started = time.perf_counter()
    result = text_coding.preview(["a" * 300_000, "b" * 200_000 + "ing"], payload)
    assert time.perf_counter() - started < 2
    assert result["coverage"]["uncoded"] == 2
    # A word as long as a term may be is read; a longer one — no word a term
    # can name — matches none.
    assert _matches("delay*", "delay" + "s" * 195)
    assert not _matches("delay*", "delay" + "s" * 196)


def test_a_long_repetitive_answer_takes_little_time_and_memory():
    """What coding an answer keeps grows with the terms that match it, not
    with its words times the terms they could begin."""
    import tracemalloc

    nothing = _cf(_theme(1, "Nothing wrong", [f"no problem{i}" for i in range(10)]))
    started = time.perf_counter()
    assert text_coding.preview(["no " * 300_000], nothing)["coverage"]["uncoded"] == 1
    assert time.perf_counter() - started < 3
    late = _cf(
        *[_theme(c, f"T{c}", [f"late x{c}y{i}" for i in range(50)]) for c in range(1, 41)],
        multiple=True,
    )
    cf = text_coding.parse(late)
    assert cf.rule_set.rules  # compiled before measuring
    text = "late " * 4_000 + "late x3y7"
    tracemalloc.start()
    started = time.perf_counter()
    try:
        assert text_coding.explain(text, cf)["codes"] == [3]
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert time.perf_counter() - started < 3
    assert peak < 40_000_000, f"{peak / 1e6:.0f} MB"


def test_many_forms_with_a_star_inside_are_looked_up_not_tried_one_by_one():
    """Two hundred themes of forms like c*ng — within the limits — against a
    few thousand different words."""
    import random

    rng = random.Random(3)
    letters = "abcdefghijklmnop"
    themes = [
        _theme(
            code,
            f"T{code}",
            [
                rng.choice(
                    [
                        f"{rng.choice(letters)}*{rng.choice(letters)}{code}x{i}",
                        f"{rng.choice(letters)}{code}x{i}*{rng.choice(letters)}",
                        f"{rng.choice(letters)}*x{code}*{i}",
                        f"*{rng.choice(letters)}{code}x{i}*",
                    ]
                )
                for i in range(250)
            ],
        )
        for code in range(1, 201)
    ]
    cf = text_coding.parse(_cf(*themes, multiple=True))
    assert cf.rule_set.rules  # compiled before measuring
    words = list(
        {"".join(rng.choice(letters) for _ in range(rng.randint(3, 10))) for _ in range(4000)}
    )
    answers = [" ".join(rng.choices(words, k=10)) for _ in range(2000)]
    # One word made of a term of each kind (its * as "zz"), and what each
    # term, read in full, says of them.
    kinds = [
        lambda t: t.count("*") == 1 and t[1] == "*",  # a*b17x3
        lambda t: t.count("*") == 1 and t[-2] == "*",  # a17x3*b
        lambda t: t.count("*") == 2 and t[1] == "*",  # a*x17*3
        lambda t: t[0] == "*",  # *a17x3*
    ]
    probes = [
        next(t for t in theme["rules"]["include"] if kind(t)).replace("*", "zz")
        for theme, kind in zip(themes[:40:10], kinds, strict=True)
    ]
    answers.append(" ".join(probes))
    started = time.perf_counter()
    result = text_coding.preview(answers, cf)
    assert time.perf_counter() - started < 8
    expected = {
        code
        for code, theme in enumerate(themes, 1)
        for term in theme["rules"]["include"]
        if any(text_rules._glob(tuple(term.split("*")), word) for word in probes)
    }
    assert len(expected) >= 4
    assert set(result["answers"][-1]["codes"]) == expected


def test_the_rules_module_reads_a_term_back_as_written():
    term, warnings = text_rules.parse_term("Staff ~2 not_rude|RUDE*")
    assert warnings == [] and term.gap == 2
    assert [a.core for a in term.second[0]] == ["rude", "rude*"]
    assert [a.negated for a in term.second[0]] == [True, False]


# ─── the flow ────────────────────────────────────────────────────────────────


DOCUMENTS = __import__("pathlib").Path(__file__).resolve().parent / "documents"
CODEFRAME = "analysis/comment.codeframe.json"
COMMENT_FRAME = {
    "schema_version": "2.0",
    "variable": "comment",
    "multiple": True,
    "themes": [
        _theme(1, "Late", ["late"], group="Delivery"),
        _theme(2, "Damaged", ["damag*"], group="Delivery"),
        _theme(3, "Rude", ["rude"]),
        _theme(9, "Nothing", ["nothing"], exclusive=True),
    ],
    "assignments": {fp("ok"): [3]},
}
COMMENTS = [
    "Delivery was late",
    "the box was damaged, and late",
    "Nothing",
    "rude driver",
    "ok",
    "",
    "Цена высокая",
    "wasn't late",
]


def _flow(nodes, edges, name="comments"):
    return {
        "schema_version": "1.0",
        "name": name,
        "nodes": [{"id": i, "type": t, "params": p} for i, t, p in nodes],
        "edges": [
            {"from": {"node": a, "port": ap}, "to": {"node": b, "port": bp}}
            for a, ap, b, bp in edges
        ],
    }


@pytest.fixture(scope="module")
def questionnaire_doc():
    from siamang.model import loads

    return loads((DOCUMENTS / "brand_awareness.questionnaire.json").read_text("utf-8"))


def test_check_flow_knows_the_theme_variable_its_codeframe_makes(questionnaire_doc):
    """Given the codeframes, the check knows the theme variable's name when
    Theme variable is left empty, that it holds lists of codes when the
    codeframe gives several themes an answer, and a codeframe the run could
    not apply."""
    from siamang.flow import check_flow

    def flow(code_params, node):
        return _flow(
            [
                ("src", "source.responses", {}),
                ("code", "prepare.text_code", {"codeframe": CODEFRAME, **code_params}),
                node,
            ],
            [("src", "data", "code", "data"), ("code", "data", node[0], "data")],
        )

    donut = flow({}, ("bar", "visualize.bar", {"variable": "comment_theme", "layout": "donut"}))
    issues = check_flow(donut, questionnaire=questionnaire_doc)
    assert [i.code for i in issues] == ["UNKNOWN_VARIABLE"]  # no codeframe, no name
    issues = check_flow(
        donut, questionnaire=questionnaire_doc, codeframes={f"./{CODEFRAME}": COMMENT_FRAME}
    )
    assert [(i.code, i.message) for i in issues] == [
        (
            "PARAM_CONFLICT",
            "bar: comment_theme allows several answers, so its shares add up to more than "
            "100 % and are not the parts of a whole: draw them as bars (Layout = grouped).",
        )
    ]
    single = {**COMMENT_FRAME, "multiple": False}
    assert check_flow(donut, questionnaire=questionnaire_doc, codeframes={CODEFRAME: single}) == []

    named = flow(
        {"into": "why"}, ("tab", "output.tabbook", {"banner": ["why"], "path": "outputs/t.xlsx"})
    )
    issues = check_flow(
        named, questionnaire=questionnaire_doc, codeframes={CODEFRAME: COMMENT_FRAME}
    )
    assert [i.code for i in issues] == ["PARAM_CONFLICT"]
    assert "why holds multiple-choice answers" in issues[0].message

    broken = {**COMMENT_FRAME, "themes": [_theme(1, "Late", ["re:late"])]}
    issues = check_flow(named, questionnaire=questionnaire_doc, codeframes={CODEFRAME: broken})
    assert [(i.code, i.node) for i in issues] == [("PARAM_INVALID", "code")]
    assert issues[0].message.startswith(
        f"Parameter 'codeframe' of code: {CODEFRAME} cannot be applied: theme 1 (Late): "
        "include term 're:late': regular expressions are not supported"
    )
    elsewhere = {**COMMENT_FRAME, "variable": "remarks", "multiple": False}
    issues = check_flow(named, questionnaire=questionnaire_doc, codeframes={CODEFRAME: elsewhere})
    assert [i.code for i in issues] == ["UNKNOWN_VARIABLE"]
    assert "codes 'remarks', which is not a variable of this questionnaire" in issues[0].message
    # A version 1 codeframe is one theme an answer, and nominal as before.
    old = {
        "schema_version": "1.0",
        "variable": "comment",
        "themes": [{"code": 1, "label": "Late"}],
        "assignments": {fp("late"): 1},
    }
    assert check_flow(donut, questionnaire=questionnaire_doc, codeframes={CODEFRAME: old}) == []
    # A Likert chart of it, as of a multiple-choice question.
    likert = flow({}, ("lik", "visualize.likert", {"items": ["comment_theme"]}))
    issues = check_flow(
        likert, questionnaire=questionnaire_doc, codeframes={CODEFRAME: COMMENT_FRAME}
    )
    assert [(i.severity, i.message) for i in issues if i.code == "PARAM_CONFLICT"] == [
        (
            "error",
            "lik: comment_theme allows several answers; a Likert chart draws items with one "
            "answer each on a scale.",
        )
    ]
    assert "PARAM_CONFLICT" not in {
        i.code
        for i in check_flow(likert, questionnaire=questionnaire_doc, codeframes={CODEFRAME: single})
    }


def test_the_check_the_run_and_the_script_agree_on_a_theme_variable_its_codeframe_names(
    questionnaire_doc, tmp_path, monkeypatch
):
    """With Theme variable empty, the name is the codeframe's: given the same
    codeframes, the check passes the flow, the run runs it and the script is
    written; without them all three refuse it alike."""
    from siamang.cli.flow import run_check, run_flow
    from siamang.flow import FlowError, FlowRunner, check_flow, generate_flow, read_codeframes
    from siamang.io import write_snapshot
    from siamang.model import from_document

    frame = {**COMMENT_FRAME, "into": "comment_theme"}
    flow = _flow(
        [
            ("src", "source.responses", {}),
            ("code", "prepare.text_code", {"codeframe": CODEFRAME}),
            ("bar", "visualize.bar", {"variable": "comment_theme"}),
        ],
        [("src", "data", "code", "data"), ("code", "data", "bar", "data")],
    )
    survey = from_document(questionnaire_doc).survey
    responses = survey.simulate(n=16, seed=4)
    responses = responses.with_frame(
        responses.frame.assign(comment=[COMMENTS[i % len(COMMENTS)] for i in range(16)])
    )
    (tmp_path / "analysis").mkdir()
    (tmp_path / CODEFRAME).write_text(json.dumps(frame), encoding="utf-8")
    codeframes = read_codeframes(flow, tmp_path)
    assert codeframes == {CODEFRAME: frame}
    assert check_flow(flow, questionnaire=questionnaire_doc, codeframes=codeframes) == []
    code = generate_flow(flow, questionnaire_doc, codeframes=codeframes)
    assert "text_coding.apply(n_src, _codeframe, into=None" in code
    runner = FlowRunner(
        flow, questionnaire=survey, questionnaire_document=questionnaire_doc, codeframes=codeframes
    )
    result = runner.run(sources={"src": responses}, cwd=tmp_path)
    assert result.ok and result.output("code", "data").frame["comment_theme"].iloc[0] == [1]
    # Without them, none of the three knows the name.
    assert [i.code for i in check_flow(flow, questionnaire=questionnaire_doc)] == [
        "UNKNOWN_VARIABLE"
    ]
    with pytest.raises(FlowError, match="unknown variable 'comment_theme'"):
        generate_flow(flow, questionnaire_doc)
    with pytest.raises(FlowError, match="unknown variable 'comment_theme'"):
        FlowRunner(flow, questionnaire=survey, questionnaire_document=questionnaire_doc)
    # The command line reads the codeframes where the flow runs.
    flow_path, q_path = tmp_path / "comments.flow.json", tmp_path / "q.json"
    flow_path.write_text(json.dumps(flow), encoding="utf-8")
    q_path.write_text(json.dumps(questionnaire_doc), encoding="utf-8")
    snapshot = write_snapshot(responses, tmp_path / "responses.csv")
    monkeypatch.chdir(tmp_path)
    assert run_check(str(flow_path), questionnaire=str(q_path)) == 0
    assert run_flow(str(flow_path), data=[str(snapshot)], questionnaire=str(q_path)) == 0
    monkeypatch.chdir(tmp_path / "analysis")
    assert run_check(str(flow_path), questionnaire=str(q_path)) == 1
    # A file that is not JSON is a codeframe the run could not apply.
    (tmp_path / CODEFRAME).write_text("{not json", encoding="utf-8")
    broken = read_codeframes(flow, tmp_path)
    issues = check_flow(flow, questionnaire=questionnaire_doc, codeframes=broken)
    assert "PARAM_INVALID" in {i.code for i in issues}


def _comments_flow():
    return _flow(
        [
            ("src", "source.responses", {}),
            ("code", "prepare.text_code", {"codeframe": CODEFRAME, "into": "comment_theme"}),
            ("bar", "visualize.bar", {"variable": "comment_theme"}),
            ("xt", "analyze.crosstab", {"row": "comment_theme", "col": "region"}),
            ("sec", "output.report_section", {"heading": "Comments"}),
            ("save", "output.save_report", {"title": "Comments", "path": "outputs/comments.md"}),
            ("exp", "output.export_file", {"path": "outputs/comments.csv"}),
        ],
        [
            ("src", "data", "code", "data"),
            ("code", "data", "bar", "data"),
            ("code", "data", "xt", "data"),
            ("code", "table", "sec", "items"),
            ("bar", "chart", "sec", "items"),
            ("xt", "table", "sec", "items"),
            ("sec", "report", "save", "sections"),
            ("code", "data", "exp", "data"),
        ],
    )


def test_a_flow_codes_by_hand_and_by_rules_and_its_script_does_the_same(
    questionnaire_doc, tmp_path
):
    """check_flow → generate_flow → FlowRunner, and the generated script run
    on its own from a snapshot writes the same report and data."""
    import os
    import subprocess
    import sys
    from pathlib import Path

    import siamang as sg
    from siamang.codegen import generate_questionnaire
    from siamang.flow import FlowRunner, check_flow, generate_flow
    from siamang.io import write_snapshot
    from siamang.model import from_document

    survey = from_document(questionnaire_doc).survey
    responses = survey.simulate(n=48, seed=4)
    responses = responses.with_frame(
        responses.frame.assign(comment=[COMMENTS[i % len(COMMENTS)] for i in range(48)])
    )
    flow = _comments_flow()
    assert (
        check_flow(flow, questionnaire=questionnaire_doc, codeframes={CODEFRAME: COMMENT_FRAME})
        == []
    )
    runner_dir, script_dir = tmp_path / "runner", tmp_path / "script"
    for root in (runner_dir, script_dir):
        (root / "analysis").mkdir(parents=True)
        (root / CODEFRAME).write_text(json.dumps(COMMENT_FRAME), encoding="utf-8")
    # Both read the same snapshot, so the export compares like with like.
    snapshot = write_snapshot(responses, script_dir / "data" / "responses.csv")
    result = FlowRunner(flow, questionnaire=survey, questionnaire_document=questionnaire_doc).run(
        sources={"src": str(snapshot)}, cwd=runner_dir
    )
    assert result.ok, [r.error for r in result.runs if r.error]
    coded = result.output("code", "data").frame["comment_theme"].tolist()
    assert coded[:8] == [[1], [1, 2], [9], [3], [3], None, None, None]
    stat = result.output("code", "stat")
    assert stat["Coded by hand"] == 6 and stat["Coded by rules"] == 24 and stat["Answered"] == 42
    report = (runner_dir / "outputs" / "comments.md").read_text("utf-8")
    assert "| Delivery (net) | 12 | 28.6 |" in report and "Coded by rules" in report

    code = generate_flow(flow, questionnaire_doc)
    assert code == generate_flow(flow, questionnaire_doc)
    assert f'_codeframe = text_coding.load("{CODEFRAME}")' in code
    (script_dir / "survey").mkdir()
    (script_dir / "survey" / "__init__.py").write_text("", encoding="utf-8")
    (script_dir / "survey" / "questionnaire.py").write_text(
        generate_questionnaire(questionnaire_doc), encoding="utf-8"
    )
    (script_dir / "scripts").mkdir()
    script = script_dir / "scripts" / "comments.py"
    script.write_text(code, encoding="utf-8")
    engine = Path(sg.__file__).resolve().parent.parent  # this checkout's engine
    env = {
        **os.environ,
        "PYTHONPATH": os.pathsep.join([str(script_dir), str(engine)]),
        "MPLBACKEND": "Agg",
    }
    completed = subprocess.run(
        [sys.executable, str(script), "--data", str(snapshot)],
        cwd=script_dir,
        env=env,
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr
    for name in ("comments.md", "comments.csv"):
        ours = (runner_dir / "outputs" / name).read_text("utf-8")
        assert ours == (script_dir / "outputs" / name).read_text("utf-8"), name

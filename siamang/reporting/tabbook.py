"""Tab book: every chosen question crossed by a banner, one Excel sheet each.

The deliverable an agency hands a client after fieldwork: a workbook with a
contents page, then one sheet per question — its answers down, a banner of
segments across (Total, then every code of each banner variable), the base
of every column, counts and percentages, and the significance letters of the
Banner table — and a notes page that says how it was all computed.

The numbers are the Banner table's (:class:`~siamang.reporting.tables.BannerTable`):
counts and column percentages come from the same helper
(:func:`siamang.data.tables._banner_pair`), and the letters from the same
two-sided z-test of column proportions within each banner variable, with the
same level, Bonferroni correction and minimum base. Two things differ, both on
purpose, and the notes say so:

* the codebook's missing codes are left out — of the question (not in its
  base), and of a banner variable (a respondent with "Refused" as region is in
  Total, not in a region column) — and counted on the sheet;
* the base of a column, for the percentages and for the test, is the
  respondents in it **who answered the question** (Kish's effective base when
  weighted); the Banner table tests on everyone in the column, which is the same
  number only when everyone answered.

A multiple-choice question (lists of codes) is one table of every option, its
base the respondents who chose at least one; its percentages add to more than
100. An interval or ratio question adds its mean and standard deviation per
column when ``means`` is on (not tested).

On weighted data the cells hold the sums of weights as they are, and the
percentages are of those sums — as the Frequencies and Crosstab tables compute
them, never of sums rounded for show; the weighted counts and bases are shown
to one decimal, as those tables show them. The workbook is written as the
engine's other workbooks are (:mod:`siamang.io.excel_text`): a label, an
answer or a banner name that begins with ``=`` stays text, never a formula
Excel would run, and a link to a sheet quotes its name.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import TYPE_CHECKING, Any

import numpy as np
import pandas as pd

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

    from siamang.data.survey_data import SurveyData

__all__ = [
    "PERCENTAGES",
    "Column",
    "QuestionTab",
    "TabBook",
    "tabulate",
    "write_tabbook",
]

#: What the percentages of a cell are of.
PERCENTAGES = ("column", "row", "none")
#: Answers a question without answer labels may have and still be tabulated;
#: more is an open answer or a number, not a list of answers.
MAX_UNLABELLED = 30
#: Excel's limit on a sheet name.
SHEET_NAME_LENGTH = 31
#: A weighted count or base is shown to one decimal, as the Frequencies and
#: Crosstab tables show it: whole, weights summing to 1 over a thousand
#: respondents would read 0 and 1 beside their percentages. The cell holds the
#: sum of weights as it is, and the percentages are of those sums.
WEIGHTED_FORMAT = "#,##0.0"

_TOTAL = "__tabbook_total__"
_ILLEGAL_SHEET = re.compile(r"[\[\]:*?/\\]")
_RESERVED_SHEETS = {"contents", "notes", "history"}


@dataclass(frozen=True)
class Column:
    """One column of the banner: Total, or one code of a banner variable."""

    #: The banner variable, None for Total.
    variable: str | None
    value: Any
    #: The code's label ("North"), "Total" for Total.
    label: str
    #: The significance letter, None for Total (it is never tested).
    letter: str | None
    #: The banner variable's label, "Total" for Total.
    block: str

    @property
    def header(self) -> str:
        """``North (A)``: the label and the letter the other columns name it by."""
        return f"{self.label} ({self.letter})" if self.letter else self.label


@dataclass
class QuestionTab:
    """One question crossed by the banner: the numbers of one sheet."""

    variable: str
    label: str
    scale: str | None
    multiple: bool
    #: ``(code, label)`` of each answer, in the codebook's order; empty when
    #: only the mean is shown.
    answers: list[tuple[Any, str]]
    columns: list[Column]
    #: Respondents who answered, per column.
    base: list[int]
    #: The sum of their weights per column (None unweighted).
    weighted_base: list[float] | None
    #: The base the test uses: Kish's effective base weighted, else ``base``.
    effective_base: list[float]
    #: ``counts[answer][column]``: weighted when the data is.
    counts: list[list[float]]
    #: Fractions (0–1) of the column's base.
    column_percent: list[list[float]]
    #: Fractions of the answer's respondents across each banner variable (Total 1).
    row_percent: list[list[float]]
    #: ``letters[answer][column]``: the columns this one is significantly higher than.
    letters: list[list[str]]
    #: ``{"Mean": [...], "Standard deviation": [...]}`` per column, when shown.
    means: dict[str, list[float]] | None = None
    #: The missing codes left out: ``[(code, count)]``.
    left_out: list[tuple[Any, int]] = field(default_factory=list)


@dataclass
class TabBook:
    """Every question tabulated, and the ones that could not be, with why."""

    tabs: list[QuestionTab]
    skipped: list[tuple[str, str]]
    columns: list[Column]
    weight: str | None
    percentages: str
    counts: bool
    letters: bool
    level: float
    correction: str
    means: bool
    min_base: int
    respondents: int
    weighted_total: float | None
    #: Per banner variable: the missing codes left out, ``[(code, count)]``.
    banner_left_out: dict[str, list[tuple[Any, int]]]
    banner_labels: list[str]

    @property
    def shows_letters(self) -> bool:
        """Whether the sheets show the letters: they compare column percentages,
        so beside a row percentage (or a bare count, which grows with its
        column) a letter would claim what the number beside it does not show."""
        return self.letters and self.percentages == "column"


def tabulate(
    data: SurveyData,
    *,
    banner: list[str],
    questions: list[str] | None = None,
    percentages: str = "column",
    counts: bool = True,
    letters: bool = True,
    level: float = 0.05,
    correction: str = "none",
    means: bool = True,
    min_base: int = 30,
) -> TabBook:
    """The numbers of a tab book, without writing it (:func:`write_tabbook`)."""

    from siamang.data import multi
    from siamang.data.inference import without_missing_codes
    from siamang.data.survey_data import SurveyData as _SurveyData
    from siamang.reporting.tables import BannerTable

    if percentages not in PERCENTAGES:
        raise ValueError(
            f"percentages must be one of {', '.join(PERCENTAGES)}; got {percentages!r}."
        )
    if correction not in ("none", "bonferroni"):
        raise ValueError("correction must be 'none' or 'bonferroni'.")
    if percentages == "none" and not counts:
        raise ValueError("With no percentages and no counts a tab book has nothing to show.")
    banner = list(dict.fromkeys(banner or []))
    if not banner:
        raise ValueError("A tab book needs at least one banner variable.")
    frame = data.frame
    if not frame.index.is_unique:
        frame = frame.reset_index(drop=True)
    for name in banner:
        if name not in frame.columns:
            raise ValueError(f"The banner variable {name!r} is not in the data.")
        if multi.is_multi(frame[name]):
            raise ValueError(
                f"{_label(data, name)} holds multiple-choice answers, and a banner column is a "
                "group of respondents that no one else is in. Explode it first (prepare.explode) "
                "and use its columns, or choose another banner variable."
            )
    weights = _weights(data, frame)

    chosen, skipped = _questions(data, frame, questions, banner)
    cleaned, left_out = without_missing_codes(frame, [*chosen, *banner], data.variables)
    cleaned = cleaned.copy()
    cleaned[_TOTAL] = 1
    clean_data = _SurveyData(frame=cleaned, variables=data.variables, weight=data.weight)
    # The Banner table's own test, on the data with its missing codes left out.
    helper = BannerTable(
        data=clean_data,
        rows=chosen or banner,
        columns=banner,
        test=letters,
        level=level,
        correction=correction,
        min_base=min_base,
    )

    columns = [Column(None, 1, "Total", None, "Total")]
    index = 0
    for name in banner:
        values = helper._values_of(name)
        if len(values) > MAX_UNLABELLED:
            raise ValueError(
                f"{_label(data, name)} has {len(values)} different values: too many columns for a "
                "banner. Band it first (prepare.bands) or choose a variable of a few groups."
            )
        labels = _value_labels(data, name)
        for value in values:
            letter = _letter(index)
            index += 1
            text = str(labels.get(value, _code_text(value)))
            columns.append(Column(name, value, text, letter, _label(data, name)))

    tabs: list[QuestionTab] = []
    for name in chosen:
        reason, tab = _tab(
            data,
            cleaned,
            name,
            columns,
            weights,
            helper,
            means=means,
            left_out=left_out.get(name, []),
        )
        if reason:
            skipped.append((name, reason))
        elif tab is not None:
            tabs.append(tab)

    return TabBook(
        tabs=tabs,
        skipped=skipped,
        columns=columns,
        weight=data.weight,
        percentages=percentages,
        counts=counts,
        letters=letters,
        level=level,
        correction=correction,
        means=means,
        min_base=min_base,
        respondents=int(len(frame)),
        weighted_total=float(weights.sum()) if weights is not None else None,
        banner_left_out={name: left_out[name] for name in banner if name in left_out},
        banner_labels=[_label(data, name) for name in banner],
    )


def write_tabbook(
    data: SurveyData,
    path: str | Path,
    *,
    banner: list[str],
    questions: list[str] | None = None,
    percentages: str = "column",
    counts: bool = True,
    letters: bool = True,
    level: float = 0.05,
    correction: str = "none",
    means: bool = True,
    min_base: int = 30,
    created: datetime | None = None,
) -> dict[str, Any]:
    """Write the tab book to ``path`` (.xlsx); return what it holds, as a stat.

    ``questions`` empty is every nominal, ordinal and multiple-choice variable
    of the codebook except the banner variables, the weight and the response
    metadata. ``created`` is the date the notes give (now, by default).
    """

    target = Path(path)
    if target.suffix.lower() != ".xlsx":
        raise ValueError(
            f"A tab book is an Excel workbook: its path must end in .xlsx (got {str(path)!r})."
        )
    book = tabulate(
        data,
        banner=banner,
        questions=questions,
        percentages=percentages,
        counts=counts,
        letters=letters,
        level=level,
        correction=correction,
        means=means,
        min_base=min_base,
    )
    target.parent.mkdir(parents=True, exist_ok=True)
    _Writer(book, data, created or datetime.now(UTC)).save(target)
    stat: dict[str, Any] = {
        "Workbook": str(path),
        "Sheets written": len(book.tabs),
        "Questions skipped": len(book.skipped),
    }
    if book.skipped:
        stat["Skipped"] = "; ".join(f"{name}: {reason}" for name, reason in book.skipped)
    stat["Banner"] = ", ".join(book.banner_labels)
    stat["Percentages"] = percentages
    stat["Test"] = _test_text(book) if book.shows_letters else _no_letters(book)
    if book.weight is not None:
        stat["Weight"] = book.weight
    return stat


# ─── the numbers ─────────────────────────────────────────────────────────────


def _label(data: SurveyData, name: str) -> str:
    if data.variables is not None and name in data.variables:
        return data.variables[name].label or name
    return name


def _value_labels(data: SurveyData, name: str) -> dict[Any, str]:
    if data.variables is not None and name in data.variables:
        return dict(data.variables[name].labels or {})
    return {}


def _code_text(code: Any) -> str:
    """A code as the codebook writes it: 1, not the 1.0 a column with a blank holds."""
    if isinstance(code, float) and code.is_integer():
        return str(int(code))
    return str(code)


def _sort_key(code: Any) -> tuple[int, float, str]:
    if isinstance(code, bool):
        return (0, float(code), "")
    if isinstance(code, int | float | np.integer | np.floating):
        return (0, float(code), "")
    return (1, 0.0, str(code))


def _letter(index: int) -> str:
    """The Banner table's letters: A–Z, then #27 …"""
    from siamang.reporting.tables import column_letter

    return column_letter(index)


def _weights(data: SurveyData, frame: pd.DataFrame) -> np.ndarray | None:
    if data.weight is None:
        return None
    if data.weight not in frame.columns:
        raise ValueError(f"Weight column '{data.weight}' not found in frame.")
    return pd.to_numeric(frame[data.weight], errors="coerce").fillna(0.0).to_numpy(dtype=float)


def _questions(
    data: SurveyData, frame: pd.DataFrame, questions: list[str] | None, banner: list[str]
) -> tuple[list[str], list[tuple[str, str]]]:
    """The questions to tabulate, and the ones named that the data does not have."""
    from siamang.data import multi
    from siamang.data.checks import METADATA_COLUMNS

    skipped: list[tuple[str, str]] = []
    if questions:
        chosen = []
        for name in dict.fromkeys(questions):
            if name not in frame.columns:
                skipped.append((name, "not in the data"))
            else:
                chosen.append(name)
        return chosen, skipped
    if not data.variables:
        raise ValueError(
            "This data has no codebook, so a tab book cannot tell its questions from its other "
            "columns. Name the Questions, or load the data with its questionnaire."
        )
    chosen = []
    for name, variable in data.variables.items():
        if (
            name in banner
            or name == data.weight
            or name in METADATA_COLUMNS
            or name.startswith("url_")
            or (variable.role or "") in {"weight", "id"}
        ):
            continue
        listed = name in frame.columns and multi.is_multi(frame[name])
        if variable.scale not in ("nominal", "ordinal") and not listed:
            continue
        if name not in frame.columns:
            skipped.append((name, "not in the data"))
            continue
        chosen.append(name)
    return chosen, skipped


def _answers(data: SurveyData, series: pd.Series, name: str, multiple: bool) -> list[Any]:
    """A question's answers: the codebook's in its order (not its missing
    codes), then the values it does not label, by code."""
    from siamang.data import multi

    variable = data.variables.get(name) if data.variables is not None else None
    missing = {_code_text(code) for code in variable.missing_values} if variable else set()
    declared = [code for code in _value_labels(data, name) if _code_text(code) not in missing]
    present = multi.codes_in(series) if multiple else series.dropna().unique().tolist()
    known = {_code_text(code) for code in declared}
    extra = sorted(
        (
            value
            for value in present
            if _code_text(value) not in known and _code_text(value) not in missing
        ),
        key=_sort_key,
    )
    return declared + extra


def _tab(
    data: SurveyData,
    cleaned: pd.DataFrame,
    name: str,
    columns: list[Column],
    weights: np.ndarray | None,
    helper: Any,
    *,
    means: bool,
    left_out: list[tuple[Any, int]],
) -> tuple[str | None, QuestionTab | None]:
    """One question's table, or why it has none."""
    from siamang.data import multi
    from siamang.data.tables import _banner_pair

    series = cleaned[name]
    variable = data.variables.get(name) if data.variables is not None else None
    multiple = multi.is_multi(series)
    answered = (
        multi.responded(series).to_numpy(dtype=bool)
        if multiple
        else series.notna().to_numpy(dtype=bool)
    )
    if not answered.any():
        return "nobody answered it", None
    scale = variable.scale if variable is not None else None
    numbers = None if multiple else pd.to_numeric(series, errors="coerce")
    numeric = numbers is not None and bool(numbers[answered].notna().all())
    if scale is None and numeric and series[answered].nunique() > MAX_UNLABELLED:
        scale = "ratio"  # a number the codebook does not describe
    answers = _answers(data, series, name, multiple)
    labelled = bool(_value_labels(data, name))
    listable = labelled or len(answers) <= MAX_UNLABELLED
    with_means = means and scale in ("interval", "ratio")
    if with_means and not numeric:
        sample = (
            series[answered & (numbers.isna().to_numpy())].iloc[0] if numbers is not None else ""
        )
        return (
            f"{scale}, but it holds text that is not a number (for example {sample!r}), so it has "
            "no mean",
            None,
        )
    if scale in ("interval", "ratio") and not listable:
        answers = []
        if not with_means:
            return (
                f"{scale} with {series[answered].nunique()} different values and no answer "
                "labels: it has no answers to count — turn on Means",
                None,
            )
    elif not listable:
        return (
            f"{len(answers)} different answers and no answer labels — an open answer? Code it "
            "first (Code open answers) or band it (Bands)",
            None,
        )

    masks = [
        np.ones(len(cleaned), dtype=bool)
        if column.variable is None
        else (cleaned[column.variable] == column.value).fillna(False).to_numpy(dtype=bool)
        for column in columns
    ]
    kept = [answered & mask for mask in masks]
    base = [int(mask.sum()) for mask in kept]
    weighted_base = [float(weights[mask].sum()) for mask in kept] if weights is not None else None
    effective = [
        _effective(weights[mask]) if weights is not None else float(mask.sum()) for mask in kept
    ]

    counts: list[list[float]] = []
    if answers and multiple:
        for code in answers:
            chose = multi.reach(series, code).fillna(False).to_numpy(dtype=bool) & answered
            counts.append(
                [
                    float(weights[chose & mask].sum())
                    if weights is not None
                    else float((chose & mask).sum())
                    for mask in masks
                ]
            )
    elif answers:
        # The Banner table's own counts, one banner variable at a time.
        found: dict[tuple[str, str | None, str], float] = {}
        blocks = [_TOTAL, *dict.fromkeys(c.variable for c in columns if c.variable)]
        weight_column = data.weight
        for block in blocks:
            pair = _banner_pair(cleaned, name, block, weight_column, data.variables, labels=False)
            key = None if block == _TOTAL else block
            for row_value, column_value, n in zip(
                pair["row_value"], pair["column_value"], pair["n"], strict=True
            ):
                found[(_code_text(row_value), key, _code_text(column_value))] = float(n)
        for code in answers:
            counts.append(
                [
                    found.get(
                        (
                            _code_text(code),
                            column.variable,
                            _code_text(column.value) if column.variable else "1",
                        ),
                        0.0,
                    )
                    for column in columns
                ]
            )

    totals = weighted_base if weighted_base is not None else [float(b) for b in base]
    column_percent = [
        [
            count / total if total > 0 else float("nan")
            for count, total in zip(row, totals, strict=True)
        ]
        for row in counts
    ]
    row_percent = [_row_percent(row, columns) for row in counts]

    tested: list[list[str]] = []
    for position, code in enumerate(answers):
        marks = [""] * len(columns)
        by_block: dict[str, list[int]] = {}
        for at, column in enumerate(columns):
            if column.variable is not None:
                by_block.setdefault(column.variable, []).append(at)
        for block_variable, positions in by_block.items():
            block = [(block_variable, columns[at].value, columns[at].header) for at in positions]
            key = _code_text(code)
            shares = {
                (key, block_variable, columns[at].value): column_percent[position][at]
                if np.isfinite(column_percent[position][at])
                else 0.0
                for at in positions
            }
            bases = {(block_variable, columns[at].value): effective[at] for at in positions}
            letters = {
                (block_variable, columns[at].value): str(columns[at].letter) for at in positions
            }
            found_marks = helper._letters_for(block, key, shares, bases, letters, np)
            for at in positions:
                marks[at] = found_marks.get((block_variable, columns[at].value), "")
        tested.append(marks)

    mean_rows = None
    if with_means and numbers is not None:
        from siamang.reporting.tables import _weighted_summary

        values = numbers.to_numpy(dtype=float)
        mean_rows = {"Mean": [], "Standard deviation": []}
        for mask in kept:
            chosen = values[mask]
            if weights is not None:
                mean, sd, _median, _n = _weighted_summary(chosen, weights[mask])
            else:
                mean = float(chosen.mean()) if len(chosen) else float("nan")
                sd = float(chosen.std(ddof=1)) if len(chosen) > 1 else float("nan")
            mean_rows["Mean"].append(mean)
            mean_rows["Standard deviation"].append(sd)

    labels = _value_labels(data, name)
    return None, QuestionTab(
        variable=name,
        label=_label(data, name),
        scale=scale,
        multiple=multiple,
        answers=[(code, str(labels.get(code, _code_text(code)))) for code in answers],
        columns=columns,
        base=base,
        weighted_base=weighted_base,
        effective_base=effective,
        counts=counts,
        column_percent=column_percent,
        row_percent=row_percent,
        letters=tested,
        means=mean_rows,
        left_out=list(left_out),
    )


def _effective(weights: np.ndarray) -> float:
    total = float(weights.sum())
    squared = float((weights**2).sum())
    return total**2 / squared if squared > 0 else 0.0


def _row_percent(row: list[float], columns: list[Column]) -> list[float]:
    """Each count as a share of the answer's respondents in its banner variable
    (who have a value of it); Total's is the whole."""
    totals: dict[str | None, float] = {}
    for count, column in zip(row, columns, strict=True):
        totals[column.variable] = totals.get(column.variable, 0.0) + count
    return [
        count / totals[column.variable] if totals[column.variable] > 0 else float("nan")
        for count, column in zip(row, columns, strict=True)
    ]


def _no_letters(book: TabBook) -> str:
    """Why a tab book shows no letters."""
    if not book.letters:
        return "not run"
    shown = "percentages of the row" if book.percentages == "row" else "counts only"
    return f"no letters: they compare column percentages, and this book shows {shown}"


def _test_text(book: TabBook) -> str:
    text = f"two-sided z-test of column proportions at {book.level:g}, within each banner variable"
    if book.correction == "bonferroni":
        text += ", Bonferroni-corrected"
    return text


# ─── the workbook ────────────────────────────────────────────────────────────


class _Writer:
    """The workbook: Contents, one sheet per question, Notes."""

    def __init__(self, book: TabBook, data: SurveyData, created: datetime) -> None:
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        self.book = book
        self.data = data
        self.created = created
        self.bold = Font(bold=True)
        self.title = Font(bold=True, size=13)
        self.heading = Font(bold=True, size=15)
        self.muted = Font(italic=True, color="595959")
        self.letter_font = Font(bold=True, color="1F4E79", size=9)
        self.link = Font(color="0563C1", underline="single")
        self.header_fill = PatternFill("solid", fgColor="F2F2F2")
        self.base_fill = PatternFill("solid", fgColor="FAFAFA")
        self.center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.wrap = Alignment(wrap_text=True, vertical="top")
        self.left = Alignment(horizontal="left", wrap_text=True, vertical="top")
        self.rule = Border(bottom=Side(style="thin", color="BFBFBF"))

    def save(self, path: Path) -> None:
        from openpyxl import Workbook

        from siamang.io.excel_text import as_text

        workbook = Workbook()
        contents = workbook.active
        contents.title = "Contents"
        names = _sheet_names([tab.variable for tab in self.book.tabs])
        for tab, name in zip(self.book.tabs, names, strict=True):
            self._question(workbook.create_sheet(name), tab)
        notes = workbook.create_sheet("Notes")
        self._contents(contents, names)
        self._notes(notes)
        # A label, an answer or a banner name that begins with "=" is text,
        # never a formula Excel would run (as every workbook the engine writes).
        as_text(workbook.worksheets)
        workbook.save(path)

    # ── cells ──
    def _put(self, sheet: Worksheet, row: int, column: int, value: Any, **style: Any) -> Any:
        if isinstance(value, np.generic):
            value = value.item()
        if isinstance(value, float) and not np.isfinite(value):
            value = None  # an undefined number is a blank cell, not "nan"
        if isinstance(value, str):
            value = _text(value)
        cell = sheet.cell(row=row, column=column, value=value)
        for key, setting in style.items():
            if setting is not None:
                setattr(cell, key, setting)
        return cell

    def _link(self, sheet: Worksheet, row: int, column: int, text: str, target: str) -> None:
        from openpyxl.worksheet.hyperlink import Hyperlink

        from siamang.io.excel_text import sheet_link

        cell = self._put(sheet, row, column, text, font=self.link)
        cell.hyperlink = Hyperlink(
            ref=cell.coordinate, location=sheet_link(target), display=_text(text)
        )

    # ── Contents ──
    def _contents(self, sheet: Worksheet, names: list[str]) -> None:
        book = self.book
        self._put(sheet, 1, 1, "Tab book", font=self.heading)
        self._put(
            sheet,
            2,
            1,
            f"{len(book.tabs)} questions × Total, {', '.join(book.banner_labels)}",
            font=self.muted,
        )
        headers = ["#", "Question", "Variable", "Base", "Sheet"]
        for column, text in enumerate(headers, start=1):
            self._put(
                sheet, 4, column, text, font=self.bold, fill=self.header_fill, border=self.rule
            )
        row = 5
        for number, (tab, name) in enumerate(zip(book.tabs, names, strict=True), start=1):
            self._put(sheet, row, 1, number)
            self._link(sheet, row, 2, tab.label, name)
            self._put(sheet, row, 3, tab.variable)
            self._put(sheet, row, 4, tab.base[0], number_format="#,##0")
            self._link(sheet, row, 5, name, name)
            row += 1
        row += 1
        self._link(sheet, row, 2, "Notes: how the tables were computed", "Notes")
        if book.skipped:
            row += 2
            self._put(sheet, row, 1, "Not tabulated", font=self.bold)
            row += 1
            for column, text in enumerate(["", "Variable", "Why"], start=1):
                if text:
                    self._put(sheet, row, column, text, font=self.bold, fill=self.header_fill)
            for name, reason in book.skipped:
                row += 1
                self._put(sheet, row, 2, _label(self.data, name))
                self._put(sheet, row, 3, reason, alignment=self.wrap)
        sheet.column_dimensions["A"].width = 6
        longest = max([len(tab.label) for tab in book.tabs] + [24])
        sheet.column_dimensions["B"].width = min(70, longest + 2)
        sheet.column_dimensions["C"].width = 28 if book.skipped else 18
        sheet.column_dimensions["D"].width = 10
        sheet.column_dimensions["E"].width = 34
        sheet.freeze_panes = "A5"

    # ── one question ──
    def _question(self, sheet: Worksheet, tab: QuestionTab) -> None:
        from openpyxl.utils import get_column_letter

        book = self.book
        letters = book.shows_letters
        weighted = tab.weighted_base is not None
        self._put(sheet, 1, 1, tab.label, font=self.title)
        kind = tab.scale or "no scale in the codebook"
        detail = f"{tab.variable} · {kind}" + (" · multiple answers" if tab.multiple else "")
        self._put(sheet, 2, 1, f"{detail} · Base: respondents who answered", font=self.muted)
        self._link(sheet, 3, 1, "← Contents", "Contents")

        # Where each banner column's numbers go: Total one column, the others
        # a value column and — with letters — a narrow one for its letters.
        spots: list[tuple[int, int | None]] = []
        at = 2
        for column in tab.columns:
            letter_at = at + 1 if letters and column.letter else None
            spots.append((at, letter_at))
            at += 2 if letter_at else 1
        last = at - 1

        # Header: the banner variable over its columns, then each column.
        head, sub = 5, 6
        blocks: dict[str, list[int]] = {}
        for column, (value_at, letter_at) in zip(tab.columns, spots, strict=True):
            blocks.setdefault(column.block, []).extend(
                [value_at] + ([letter_at] if letter_at else [])
            )
        for column, (value_at, letter_at) in zip(tab.columns, spots, strict=True):
            # Without letters on the sheet a column needs no letter to be named by.
            title = column.header if letters else column.label
            header = dict(font=self.bold, alignment=self.center, fill=self.header_fill)
            if column.variable is None:
                # Total: one header over both rows.
                self._put(sheet, head, value_at, column.header, **header)
                self._put(sheet, sub, value_at, None, fill=self.header_fill, border=self.rule)
                sheet.merge_cells(
                    start_row=head, start_column=value_at, end_row=sub, end_column=value_at
                )
                continue
            self._put(sheet, sub, value_at, title, border=self.rule, **header)
            if letter_at:
                self._put(sheet, sub, letter_at, None, fill=self.header_fill, border=self.rule)
        for block, used in blocks.items():
            if block == "Total" and len(used) == 1:
                continue
            first, end = min(used), max(used)
            self._put(
                sheet,
                head,
                first,
                block,
                font=self.bold,
                alignment=self.center,
                fill=self.header_fill,
            )
            if end > first:
                sheet.merge_cells(start_row=head, start_column=first, end_row=head, end_column=end)
        self._put(sheet, sub, 1, None, fill=self.header_fill, border=self.rule)
        self._put(sheet, head, 1, None, fill=self.header_fill)

        row = sub + 1
        self._put(
            sheet,
            row,
            1,
            "Base (unweighted)" if weighted else "Base",
            font=self.bold,
            fill=self.base_fill,
        )
        for value, (value_at, _letter) in zip(tab.base, spots, strict=True):
            self._put(
                sheet,
                row,
                value_at,
                value,
                number_format="#,##0",
                font=self.bold,
                fill=self.base_fill,
            )
        if weighted:
            row += 1
            self._put(sheet, row, 1, "Base (weighted)", font=self.bold, fill=self.base_fill)
            for value, (value_at, _letter) in zip(tab.weighted_base or [], spots, strict=True):
                self._put(
                    sheet,
                    row,
                    value_at,
                    value,
                    number_format=WEIGHTED_FORMAT,
                    font=self.bold,
                    fill=self.base_fill,
                )

        # The headers and the bases stay in view over a long list of answers.
        frozen = row + 1
        percent = {"column": tab.column_percent, "row": tab.row_percent}.get(book.percentages)
        for position, (_code, label) in enumerate(tab.answers):
            lines: list[tuple[list[float], str]] = []
            if book.counts:
                lines.append((tab.counts[position], WEIGHTED_FORMAT if weighted else "#,##0"))
            if percent is not None:
                lines.append((percent[position], "0.0%"))
            # The letters compare column percentages: beside them.
            marked = len(lines) - 1
            for which, (cells, number_format) in enumerate(lines):
                row += 1
                if which == 0:
                    self._put(sheet, row, 1, label, alignment=self.wrap)
                for at, (value, (value_at, letter_at)) in enumerate(zip(cells, spots, strict=True)):
                    self._put(sheet, row, value_at, value, number_format=number_format)
                    mark = tab.letters[position][at]
                    if letter_at and which == marked and mark:
                        self._put(sheet, row, letter_at, mark, font=self.letter_font)
            for column in range(1, last + 1):
                sheet.cell(row=row, column=column).border = self.rule

        if tab.means:
            for name, values in tab.means.items():
                row += 1
                self._put(sheet, row, 1, name, font=self.bold if name == "Mean" else None)
                for value, (value_at, _letter) in zip(values, spots, strict=True):
                    self._put(sheet, row, value_at, value, number_format="0.00")

        row += 2
        for line in self._footnotes(tab):
            self._put(sheet, row, 1, line, font=self.muted)
            row += 1

        width = min(
            60,
            max([len(label) for _code, label in tab.answers] + [len("Standard deviation"), 18]) + 2,
        )
        sheet.column_dimensions["A"].width = width
        for column, (value_at, letter_at) in zip(tab.columns, spots, strict=True):
            # Wide enough for the header on one line up to 22; longer wraps.
            title = column.header if letters else column.label
            sheet.column_dimensions[get_column_letter(value_at)].width = max(
                10, min(22, len(title) + 3)
            )
            if letter_at:
                sheet.column_dimensions[get_column_letter(letter_at)].width = 5
        sheet.row_dimensions[sub].height = 30
        sheet.freeze_panes = sheet.cell(row=frozen, column=2)

    def _footnotes(self, tab: QuestionTab) -> list[str]:
        book = self.book
        lines = []
        if book.percentages == "column":
            lines.append("Percentages are of the column's respondents who answered.")
        elif book.percentages == "row":
            lines.append(
                "Percentages are of the row: the answer's respondents across each banner variable."
            )
        if tab.multiple:
            lines.append("Several answers were allowed, so the percentages add to more than 100%.")
        if book.letters and tab.answers and not book.shows_letters:
            lines.append(f"{_no_letters(book).capitalize()}.")
        elif book.letters and tab.answers:
            lines.append(
                "Letters: the column percentage is significantly higher than in the column of "
                f"that letter ({_test_text(book)}; a column with a base below {book.min_base} "
                "is not tested)."
            )
        if tab.means:
            lines.append("Means are not tested.")
        if tab.left_out:
            from siamang.data.inference import missing_codes_note

            note = missing_codes_note({tab.variable: tab.left_out}, self.data.variables)
            lines.append(f"Left out, as missing codes: {note}.")
        if book.weight is not None:
            lines.append(
                f"Weighted by '{book.weight}'; the unweighted base is the number of respondents."
            )
        return lines

    # ── Notes ──
    def _notes(self, sheet: Worksheet) -> None:
        from siamang.data.inference import missing_codes_note

        book = self.book
        self._put(sheet, 1, 1, "Notes", font=self.heading)
        self._link(sheet, 2, 1, "← Contents", "Contents")
        banner = []
        for label in book.banner_labels:
            letters = [
                c.letter
                for c in book.columns
                if c.block == label and c.letter and book.shows_letters
            ]
            span = (
                f" ({letters[0]}–{letters[-1]})"
                if len(letters) > 1
                else (f" ({letters[0]})" if letters else "")
            )
            banner.append(f"{label}{span}")
        missing = [
            missing_codes_note({tab.variable: tab.left_out}, self.data.variables)
            for tab in book.tabs
            if tab.left_out
        ]
        missing += [
            missing_codes_note({name: found}, self.data.variables)
            + " — in Total, in no column of the banner"
            for name, found in book.banner_left_out.items()
        ]
        rows: list[tuple[str, Any]] = [
            ("Created", self.created.strftime("%Y-%m-%d %H:%M UTC")),
            ("Respondents", book.respondents),
            ("Weight", book.weight or "none: every count is of respondents"),
        ]
        if book.weighted_total is not None:
            rows.append(("Weighted total", round(book.weighted_total, 1)))
        rows += [
            ("Banner", "Total, " + ", ".join(banner)),
            (
                "Percentages",
                {
                    "column": "of the column: its respondents who answered the question",
                    "row": "of the row: the answer's respondents across each banner variable",
                    "none": "none: counts only",
                }[book.percentages],
            ),
            ("Counts", "shown" if book.counts else "not shown"),
            (
                "Base",
                "respondents who answered the question in the column; for a multiple-choice "
                "question, those who chose at least one answer",
            ),
            (
                "Test",
                _test_text(book)
                + ": each column is compared with the other columns of its banner variable only "
                "(their respondents are different people); Total is not tested"
                if book.shows_letters
                else _no_letters(book),
            ),
        ]
        if book.shows_letters:
            rows += [
                ("Alpha", book.level),
                (
                    "Bonferroni",
                    "yes: alpha divided by the pairs compared within each banner variable"
                    if book.correction == "bonferroni"
                    else "no",
                ),
                (
                    "Minimum base for a test",
                    f"{book.min_base} respondents"
                    + (" (Kish's effective base)" if book.weight is not None else ""),
                ),
                (
                    "Letters",
                    "a letter beside a cell names a column of the same banner variable whose "
                    "column percentage this one's is significantly higher than",
                ),
            ]
        rows.append(
            (
                "Means",
                "mean and standard deviation of interval and ratio questions, not tested"
                if book.means
                else "not shown",
            )
        )
        rows.append(("Missing codes left out", "\n".join(missing) if missing else "none met"))
        if book.skipped:
            rows.append(
                (
                    "Not tabulated",
                    "\n".join(f"{name}: {reason}" for name, reason in book.skipped),
                )
            )
        for row, (key, value) in enumerate(rows, start=4):
            self._put(sheet, row, 1, key, font=self.bold, alignment=self.wrap)
            # A number among sentences reads from the left like them.
            self._put(sheet, row, 2, value, alignment=self.left)
        sheet.column_dimensions["A"].width = 26
        sheet.column_dimensions["B"].width = 100


def _text(value: str) -> str:
    """Text a cell can hold: no control characters XML refuses."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    return ILLEGAL_CHARACTERS_RE.sub("", value)


def _sheet_names(variables: list[str]) -> list[str]:
    """A sheet name per question: its variable name, without the characters
    Excel refuses, at most 31 long, unique ignoring case (as Excel compares)."""
    taken = set(_RESERVED_SHEETS)
    names = []
    for variable in variables:
        stem = _ILLEGAL_SHEET.sub("_", _text(variable)).strip("'") or "Question"
        # Excel refuses a name that begins or ends with an apostrophe, as a cut may leave it.
        stem = stem[:SHEET_NAME_LENGTH].rstrip("'")
        name, number = stem, 1
        while name.lower() in taken:
            number += 1
            suffix = f"~{number}"
            name = stem[: SHEET_NAME_LENGTH - len(suffix)].rstrip("'") + suffix
        taken.add(name.lower())
        names.append(name)
    return names

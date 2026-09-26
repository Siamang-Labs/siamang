"""Every table of a report in one Excel workbook.

A report's reader who wants to work with its numbers asks for them in Excel,
and copying each table out of an HTML page loses its significance letters,
its post-hoc pairs and its notes. :func:`save_tables` writes one sheet per
table of a :class:`~siamang.reporting.document.Report`:

- each sheet is what the table's own ``export_xlsx`` writes — a Banner with
  its significance letters, Group means with its post-hoc pairs on a sheet of
  their own beside it — so the workbook and a table exported alone can never
  differ; a bare DataFrame is written as the report prints it (no index);
- the statistics a table prints under itself (χ², p, N, the weight, the
  missing codes left out) are written under it, after an empty row — and the
  post-hoc pairs' own (the method, which way a difference runs, that p is
  adjusted) under the pairs;
- a sheet is named by the table's caption, else by its section's heading,
  else by the variable it describes, else ``Table <n>`` — cut to Excel's 31
  characters, without the characters Excel refuses, and made unique;
- the first sheet, ``Contents``, lists every sheet with its section and full
  caption, each a link to the sheet;
- text is text: a respondent's answer, a label or a caption that begins with
  ``=`` is written as a string, never as a formula Excel would run
  (:mod:`siamang.io.excel_text`).

Charts, text and statistics lines are not tables and are left out.
"""

from __future__ import annotations

import re
import tempfile
from copy import copy
from pathlib import Path
from typing import TYPE_CHECKING, Any

import numpy as np
import pandas as pd

if TYPE_CHECKING:
    from siamang.reporting.document import Report

#: What Excel refuses in a sheet name, and its longest.
_REFUSED = re.compile(r"[\[\]:*?/\\]")
SHEET_NAME_LENGTH = 31
CONTENTS = "Contents"
_HEADING = re.compile(r"^#{1,6}\s+(.+?)\s*#*\s*$")


def save_tables(report: Report, path: str | Path) -> Path:
    """Write every table of ``report`` to the workbook ``path`` (``.xlsx``)."""

    import openpyxl
    from openpyxl.styles import Font

    from siamang.io.excel_text import as_text

    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"The tables are written to an .xlsx workbook; got {path.name!r}.")
    path.parent.mkdir(parents=True, exist_ok=True)
    book = openpyxl.Workbook()
    contents = book.active
    contents.title = CONTENTS
    names = SheetNames(reserved=[CONTENTS])
    entries: list[tuple[str, str, str]] = []
    heading: str | None = None
    number = 0
    for kind, payload in report._blocks:
        if kind == "md":
            heading = _heading(str(payload)) or heading
            continue
        if kind != "table":
            continue
        component, caption = payload[0], payload[1]  # type: ignore[index]
        number += 1
        base = caption or heading or _own_name(component) or f"Table {number}"
        for index, (own, sheet) in enumerate(_sheets(component)):
            name = names.take(base, suffix=own if index else None)
            target = book.create_sheet(name)
            _copy(sheet, target)
            _write_stats(target, _sheet_source(component, index, own))
            _fit_columns(target)
            described = caption or _described(component)
            entries.append(
                (name, heading or "", described if not index else f"{described} — {own}")
            )

    contents["A1"] = report.title or "Tables"
    contents["A1"].font = Font(bold=True, size=14)
    contents["A2"] = (
        "One sheet per table of the report, its statistics under it; charts are not included."
    )
    if entries:
        for column, title in enumerate(("Sheet", "Section", "Table"), 1):
            contents.cell(row=4, column=column, value=title).font = Font(bold=True)
        for row, (name, section, described) in enumerate(entries, 5):
            link = contents.cell(row=row, column=1, value=name)
            link.hyperlink = f"#'{name}'!A1"
            link.style = "Hyperlink"
            contents.cell(row=row, column=2, value=section or None)
            contents.cell(row=row, column=3, value=described)
    else:
        contents["A4"] = "This report has no tables."
    _fit_columns(contents, first_row=4)
    as_text(book.worksheets)
    book.save(path)
    return path


class SheetNames:
    """Sheet names Excel accepts: at most 31 characters, none of ``[]:*?/\\``,
    no apostrophe at either end, and unique regardless of case."""

    def __init__(self, reserved: list[str] | None = None) -> None:
        self._used = {name.casefold() for name in reserved or []}
        # Excel keeps "History" for itself.
        self._used.add("history")

    def take(self, base: str, suffix: str | None = None) -> str:
        stem = _clean(base) or "Table"
        tail = f" – {_clean(suffix)}" if suffix else ""
        count = 1
        while True:
            mark = f" ({count})" if count > 1 else ""
            room = SHEET_NAME_LENGTH - len(tail) - len(mark)
            name = (stem[: max(room, 1)].rstrip(" '") + tail + mark)[:SHEET_NAME_LENGTH]
            name = name.strip("'") or "Table"
            if name.casefold() not in self._used:
                self._used.add(name.casefold())
                return name
            count += 1


def _clean(text: Any) -> str:
    return " ".join(_REFUSED.sub(" ", str(text)).split()).strip("'")


def _heading(markdown: str) -> str | None:
    """A block's heading (``## Satisfaction``), if it starts with one."""

    lines = markdown.strip().splitlines()
    match = _HEADING.match(lines[0]) if lines else None
    return match.group(1).strip() if match else None


def _own_name(component: Any) -> str | None:
    """What a table says it describes (a Frequencies table's variable)."""

    stats = getattr(component, "stats", None)
    if isinstance(stats, dict) and isinstance(stats.get("Variable"), str):
        return stats["Variable"]
    return None


def _analysis_name(component: Any) -> str | None:
    """What a table of the later analyses describes, and which of the
    analysis's tables it is: ``Region × Gender — rows (Region)`` of a
    Perceptual map, ``Gabor-Granger — curves`` of Price sensitivity, the
    outcome of Key drivers, the test of Paired tests."""

    stats = getattr(component, "stats", None)
    stats = stats if isinstance(stats, dict) else {}
    analysis = getattr(component, "analysis", None)
    kind = type(component).__name__
    if kind == "MapTable" and analysis is not None:
        title = f"{analysis.row_title} × {analysis.column_title}"
        part = (
            "inertia"
            if component is analysis.table
            else f"rows ({analysis.row_title})"
            if component is analysis.rows
            else f"columns ({analysis.column_title})"
            if component is analysis.columns
            else None
        )
        return f"{title} — {part}" if part else title
    if kind == "PriceTable" and analysis is not None:
        method = str(stats.get("Method") or analysis.method)
        if component is analysis.curves:
            return f"{method} — curves"
        points = "price points" if analysis.method == "van_westendorp" else "demand and revenue"
        return f"{method} — {points}"
    if kind == "DriverTable" and isinstance(stats.get("Outcome"), str):
        return str(stats["Outcome"])
    if isinstance(stats.get("Test"), str):
        return str(stats["Test"])
    return None


#: What each table component is, for the contents of a table without a caption.
_KINDS = {
    "FreqTable": "Frequencies",
    "CrossTable": "Crosstab",
    "GroupMeanTable": "Group means",
    "BannerTable": "Banner table",
    "NpsTable": "Net Promoter Score",
    "DescriptivesTable": "Descriptive statistics",
    "DataCheckTable": "Data check",
    "TTestTable": "t-test",
    "CorrelationMatrixTable": "Correlation matrix",
    "PostHocTable": "Post-hoc comparisons",
    "QualityTable": "Response quality",
    "ThemeTable": "Coded open answers",
    "MaxDiffTable": "MaxDiff",
    "ConjointTable": "Conjoint",
    "ShareTable": "Share of preference",
    "DriverTable": "Key drivers",
    "MapTable": "Perceptual map",
    "PriceTable": "Price sensitivity",
}


def _described(component: Any) -> str:
    """``"Frequencies: Region"``: the kind of table and what it describes."""

    name = type(component).__name__
    own = _own_name(component) or _analysis_name(component)
    kind = _KINDS.get(name, "Table")
    if name == "ResultTable" and own and own == getattr(component, "stats", {}).get("Test"):
        kind = "Paired tests"  # its own name is the test: "Paired tests: Cochran's Q"
    return f"{kind}: {own}" if own else kind


def _sheet_source(component: Any, index: int, own: str) -> Any:
    """The table whose statistics go under sheet ``index`` of ``component``'s
    export: the table itself on its first sheet, its post-hoc pairs' on theirs."""

    if index == 0:
        return component
    posthoc = getattr(component, "posthoc_table", None)
    return posthoc if own == "Post-hoc" and posthoc is not None else None


def _sheets(component: Any) -> list[tuple[str, Any]]:
    """The sheets the table's own export writes, as (its sheet name, worksheet)."""

    import openpyxl

    from siamang.reporting.tables import SurveyTable

    with tempfile.TemporaryDirectory() as directory:
        file = Path(directory) / "table.xlsx"
        if isinstance(component, SurveyTable):
            component.export_xlsx(file)
        else:
            _printable(component).to_excel(file, index=False, sheet_name="Table")
        book = openpyxl.load_workbook(file)
    return [(sheet.title, sheet) for sheet in book.worksheets]


def _printable(frame: pd.DataFrame) -> pd.DataFrame:
    """A bare DataFrame as the report prints it: no index, one header row."""

    if isinstance(frame.columns, pd.MultiIndex):
        frame = frame.copy()
        frame.columns = [
            " / ".join(str(part) for part in column if str(part)) for column in frame.columns
        ]
    return frame


def _copy(source: Any, target: Any) -> None:
    for row in source.iter_rows():
        for cell in row:
            written = target.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                written.font = copy(cell.font)
                written.border = copy(cell.border)
                written.fill = copy(cell.fill)
                written.alignment = copy(cell.alignment)
                written.protection = copy(cell.protection)
                written.number_format = cell.number_format
    for merged in source.merged_cells.ranges:
        target.merge_cells(str(merged))
    for key, dimension in source.column_dimensions.items():
        if dimension.width:
            target.column_dimensions[key].width = dimension.width
    target.freeze_panes = source.freeze_panes


def _write_stats(target: Any, component: Any) -> None:
    """The statistics a table prints under itself, one per row, under it."""

    from openpyxl.styles import Font

    stats = getattr(component, "stats", None) if component is not None else None
    if not isinstance(stats, dict) or not stats:
        return
    row = target.max_row + 2
    for key, value in stats.items():
        target.cell(row=row, column=1, value=str(key)).font = Font(italic=True)
        target.cell(row=row, column=2, value=_cell(value))
        row += 1


def _cell(value: Any) -> Any:
    """A statistic as a cell holds it: a number stays a number."""

    if value is None or isinstance(value, bool | np.bool_):
        return None if value is None else str(value)
    if isinstance(value, int | np.integer):
        return int(value)
    if isinstance(value, float | np.floating):
        return float(value) if np.isfinite(value) else str(value)
    return str(value)


def _fit_columns(sheet: Any, first_row: int = 1) -> None:
    """Columns as wide as their longest value (8 to 60 characters), where the
    table's own export set no width."""

    from openpyxl.utils import get_column_letter

    widths: dict[int, int] = {}
    for row in sheet.iter_rows(min_row=first_row):
        for cell in row:
            if cell.value is not None:
                longest = max(len(line) for line in str(cell.value).split("\n"))
                widths[cell.column] = max(widths.get(cell.column, 0), longest)
    for column, width in widths.items():
        letter = get_column_letter(column)
        if not sheet.column_dimensions[letter].width:
            sheet.column_dimensions[letter].width = min(60, max(8, width + 2))


__all__ = ["CONTENTS", "SHEET_NAME_LENGTH", "SheetNames", "save_tables"]
